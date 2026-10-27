import urllib.request
import base64
import zlib
from docx import Document
from docx.shared import Inches
import shutil
import os

# 1. Generate Use Case Diagram via PlantUML (Kroki)
plantuml_code = """
@startuml
skinparam packageStyle rectangle
skinparam defaultFontSize 14
left to right direction

actor "Member Actor" as Member
actor "Admin Actor" as Admin

rectangle "Behavioral Model - Use Case Design & Actor Roles" {
  usecase "UC-1: Register & Login" as UC1
  usecase "UC-2: Upload KYC Document" as UC2
  usecase "UC-3: Explore & Request to Join plan" as UC3
  usecase "UC-4: View Unbooked Due Schedule" as UC4
  usecase "UC-5: Pay monthly Net Due via UPI" as UC5
  usecase "UC-6: Download Payment Receipt" as UC6
  
  usecase "UC-7: Review & Approve member KYC" as UC7
  usecase "UC-8: Create & Launch Chit Schemes" as UC8
  usecase "UC-9: Conduct monthly Auction Draw" as UC9
  usecase "UC-10: AI urgency Priority (Hospital/Edu)" as UC10
  usecase "UC-11: Generate Audit & Financial Reports" as UC11
}

Member --> UC1
Member --> UC2
Member --> UC3
Member --> UC4
Member --> UC5
Member --> UC6

UC1 <-- Admin
UC7 <-- Admin
UC8 <-- Admin
UC9 <-- Admin
UC10 <-- Admin
UC11 <-- Admin
@enduml
"""

def generate_kroki_png(code, diagram_type, output_filename):
    compressed = zlib.compress(code.encode('utf-8'), 9)
    b64 = base64.urlsafe_b64encode(compressed).decode('ascii')
    url = f"https://kroki.io/{diagram_type}/png/{b64}"
    req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
    try:
        with urllib.request.urlopen(req) as response:
            with open(output_filename, 'wb') as f:
                f.write(response.read())
            print(f"Saved {output_filename}")
    except Exception as e:
        print(f"Error saving {output_filename}: {e}")

use_case_path = r"C:\chit fund\use_case_diagram.png"
generate_kroki_png(plantuml_code, "plantuml", use_case_path)

# 2. Copy the user's handdrawn image
user_img_path = r"C:\Users\nares\.gemini\antigravity\brain\a8276421-67ca-4e5a-99db-e38e66b25de5\.user_uploaded\media_1791035629735.png"
target_img_path = r"C:\chit fund\handdrawn_usecase.png"
try:
    if os.path.exists(user_img_path):
        shutil.copy(user_img_path, target_img_path)
        print("Copied user image.")
    else:
        print("User image not found at path.")
except Exception as e:
    print("Error copying user image:", e)

# 3. Create the Word Document
doc = Document()

# Adjust margins to fit large images
sections = doc.sections
for section in sections:
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)
    section.top_margin = Inches(0.5)
    section.bottom_margin = Inches(0.5)

doc.add_heading("Secure Chit System - System Diagrams", 0)

doc.add_heading("1. Context Level Diagram (Level 0 DFD)", level=1)
try:
    doc.add_picture(r"C:\chit fund\dfd_level_0.png", width=Inches(7.2))
except Exception as e:
    print(f"Failed to add DFD 0: {e}")

doc.add_page_break()
doc.add_heading("2. System Overview & Subsystem Decomposition (Level 1 DFD)", level=1)
try:
    doc.add_picture(r"C:\chit fund\dfd_level_1.png", width=Inches(7.2))
except Exception as e:
    print(f"Failed to add DFD 1: {e}")

doc.add_page_break()
doc.add_heading("3. Auction Bidding & Workflow (Level 2 DFD)", level=1)
try:
    doc.add_picture(r"C:\chit fund\dfd_level_2.png", width=Inches(7.2))
except Exception as e:
    print(f"Failed to add DFD 2: {e}")

doc.add_page_break()
doc.add_heading("4. Behavioral Model - Use Case Design (Digital Version)", level=1)
try:
    doc.add_picture(use_case_path, width=Inches(7.2))
except Exception as e:
    print(f"Failed to add digital Use Case: {e}")

doc.add_page_break()
doc.add_heading("5. Original Hand-Drawn Use Case Sketch", level=1)
try:
    if os.path.exists(target_img_path):
        # Rotate image or just add it normally, word will handle it
        doc.add_picture(target_img_path, width=Inches(6.5))
except Exception as e:
    print("Could not add handdrawn image to doc:", e)

output_doc = r"C:\chit fund\SecureChit_All_Diagrams.docx"
doc.save(output_doc)
print(f"\nSuccessfully created Word Doc: {output_doc}")
