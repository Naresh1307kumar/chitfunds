import shutil
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

src = r'C:\chit fund\SecureChit_Report_With_DFDs_v5.docx'
out = r'C:\chit fund\SecureChit_Report_With_DFDs_v6.docx'
doc = Document(src)

parent = doc.element.body

def ins(elem):
    parent.append(elem)

def heading(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(13)
    ins(p._element)

def body(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.font.size = Pt(11)
    ins(p._element)

def blank():
    p = doc.add_paragraph()
    p.add_run("").font.size = Pt(6)
    ins(p._element)

# ─── 6. CONCLUSION ────────────────────────────────────────────────────────────
blank()
heading("6. CONCLUSION")
blank()

body(
    "SecureChit is a fully functional, enterprise-grade digital chit fund management platform "
    "developed to eliminate the inefficiencies, opacity, and financial risks inherent in "
    "conventional paper-based chit fund operations. By combining a modern technology stack — "
    "React 18, FastAPI 0.138, and PostgreSQL 16 — with robust domain-specific financial logic, "
    "the platform delivers a transparent, legally structured, and auditable chit fund ecosystem "
    "for both members and administrators. The reverse-auction engine with the AI Emergency "
    "Priority framework ensures fair prize allocation, the automated dividend distribution "
    "enforces the statutory 5% commission cap, and the Digital Promissory Note module with "
    "Canvas digital signature eliminates all manual paperwork. All 25 system test cases returned "
    "a PASS status, confirming that the platform meets its functional, security, and data "
    "integrity requirements."
)
blank()

body(
    "In conclusion, SecureChit demonstrates that the century-old chit fund — one of India's most "
    "widely used community savings instruments — can be safely and effectively digitized without "
    "compromising its fundamental trust-based model. The platform stands as a reliable, scalable, "
    "and legally compliant foundation that can be adopted by registered chit fund companies of "
    "any size, bringing the benefits of digital finance to millions of subscribers across India. "
    "Future enhancements such as a React Native mobile application, NACH auto-debit integration, "
    "blockchain audit trail, and AI-based member credit scoring will further strengthen the "
    "platform's reach, security, and operational efficiency."
)
blank()

doc.save(out)
print(f"Saved: {out}")

shutil.copy(out, r'C:\Users\nares\.gemini\antigravity\brain\f74afe5d-2f73-4013-bc1e-3562e05df0f5\SecureChit_Report_With_DFDs.docx')
print("Done!")
