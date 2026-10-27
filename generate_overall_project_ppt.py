import sys
import os

try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
except ImportError:
    os.system("pip install python-pptx")
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
    from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR

def build_overall_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Professional Color Palette
    NAVY = RGBColor(26, 43, 73)          # #1A2B49 (Corporate Navy)
    GOLD = RGBColor(217, 119, 6)         # Amber-600
    SLATE_DARK = RGBColor(15, 23, 42)    # Slate-900
    LIGHT_BLUE = RGBColor(239, 246, 255) # Blue-50
    BORDER_BLUE = RGBColor(59, 130, 246) # Blue-500
    WHITE = RGBColor(255, 255, 255)
    GREEN = RGBColor(16, 185, 129)       # Emerald-500
    GREEN_BG = RGBColor(236, 253, 245)    # Emerald-50
    AMBER_BG = RGBColor(254, 243, 199)   # Amber-100
    SLATE_BG = RGBColor(241, 245, 249)   # Slate-100

    def add_header(slide, title_text, category_text="SECURECHIT OVERALL PROJECT PRESENTATION"):
        tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(1.1))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p0 = tf.paragraphs[0]
        p0.text = category_text.upper()
        p0.font.size = Pt(12)
        p0.font.bold = True
        p0.font.color.rgb = GOLD
        
        p1 = tf.add_paragraph()
        p1.text = title_text
        p1.font.size = Pt(26)
        p1.font.bold = True
        p1.font.color.rgb = NAVY

    def add_card(slide, left, top, width, height, title, items, bg_color=LIGHT_BLUE, border_color=BORDER_BLUE, title_color=NAVY, font_size=13):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        shape.line.color.rgb = border_color
        shape.line.width = Pt(2.0)

        tf = shape.text_frame
        tf.word_wrap = True
        tf.vertical_anchor = MSO_ANCHOR.TOP
        tf.margin_left = Inches(0.25)
        tf.margin_right = Inches(0.25)
        tf.margin_top = Inches(0.25)
        tf.margin_bottom = Inches(0.25)

        p0 = tf.paragraphs[0]
        p0.text = title
        p0.font.size = Pt(16)
        p0.font.bold = True
        p0.font.color.rgb = title_color
        p0.space_after = Pt(8)

        for item in items:
            p = tf.add_paragraph()
            p.text = f"• {item}" if not item.startswith("  ") else f"   {item}"
            p.font.size = Pt(font_size)
            p.font.color.rgb = SLATE_DARK
            p.space_after = Pt(6)
        return shape

    # =========================================================================
    # SLIDE 1: Title Slide
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = NAVY
    bg1.line.color.rgb = NAVY

    tb1 = s1.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(11.333), Inches(4.2))
    tf1 = tb1.text_frame
    tf1.word_wrap = True

    p0 = tf1.paragraphs[0]
    p0.text = "SECURE CHIT FUND MANAGEMENT SYSTEM"
    p0.font.size = Pt(38)
    p0.font.bold = True
    p0.font.color.rgb = WHITE

    p1 = tf1.add_paragraph()
    p1.text = "Comprehensive Full-Stack Digital Savings & Auction Chit Fund Platform"
    p1.font.size = Pt(20)
    p1.font.bold = True
    p1.font.color.rgb = GOLD

    p2 = tf1.add_paragraph()
    p2.text = "\nTech Stack: FastAPI (Python 3.11) | React.js 18 (Vite) | PostgreSQL | TailwindCSS | Razorpay Gateway"
    p2.font.size = Pt(15)
    p2.font.color.rgb = WHITE

    # =========================================================================
    # SLIDE 2: Introduction
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    add_header(s2, "Introduction & Project Background", "1. PROJECT OVERVIEW")
    add_card(s2, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2), "What is a Chit Fund?", [
        "Traditional Indian financial institution combining savings and credit.",
        "Group of members contribute a fixed sum monthly into a central pool.",
        "Monthly pool is awarded to a member via auction bidding or lucky draw.",
        "Remaining discount amount is distributed as dividends among all members."
    ], bg_color=LIGHT_BLUE, border_color=BORDER_BLUE)
    add_card(s2, Inches(6.8), Inches(1.6), Inches(5.6), Inches(5.2), "Need for Digitalization", [
        "Eliminates manual ledger book-keeping & physical cash collection errors.",
        "Provides 100% transparency in monthly dividend & payout calculations.",
        "Ensures cryptographic security, digital KYC, and role-based access control.",
        "Enables instant online monthly dues payment via Razorpay / UPI."
    ], bg_color=GREEN_BG, border_color=GREEN)

    # =========================================================================
    # SLIDE 3: Existing System
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    add_header(s3, "Existing System & Limitations", "2. PROBLEM STATEMENT")
    add_card(s3, Inches(0.8), Inches(1.6), Inches(3.6), Inches(5.2), "Manual Book-Keeping", [
        "Paper ledgers & registers.",
        "Prone to record tampering.",
        "High risk of data loss.",
        "No real-time tracking for members."
    ], bg_color=AMBER_BG, border_color=GOLD)
    add_card(s3, Inches(4.8), Inches(1.6), Inches(3.6), Inches(5.2), "Lack of Transparency", [
        "Hidden company commissions.",
        "Unclear dividend calculations.",
        "Manual auction bidding disputes.",
        "No verifiable audit logs."
    ], bg_color=AMBER_BG, border_color=GOLD)
    add_card(s3, Inches(8.8), Inches(1.6), Inches(3.6), Inches(5.2), "Payment & Security Risks", [
        "Physical cash handling risks.",
        "Delayed manual verifications.",
        "Unverified member identities.",
        "No digital payment proof."
    ], bg_color=AMBER_BG, border_color=GOLD)

    # =========================================================================
    # SLIDE 4: Proposed System
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    add_header(s4, "Proposed System — SecureChit Solution", "3. SOLUTION ARCHITECTURE")
    add_card(s4, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2), "Key Improvements & Capabilities", [
        "Automated Chit Scheme Engine (Bronze, Silver, Gold).",
        "Transparent Month-by-Month Auction & Dividend Schedule Table.",
        "5% Foreman Commission Auto-Calculation & Net Payable Dues.",
        "Instant UPI / QR Code Payment Gateway Integration.",
        "Verified Digital Receipts (`REC-2026-XXXX`) with PDF printing."
    ], bg_color=GREEN_BG, border_color=GREEN)
    add_card(s4, Inches(6.8), Inches(1.6), Inches(5.6), Inches(5.2), "AI Emergency Prioritization", [
        "Rank 1: 🏥 Hospital / Medical Emergency",
        "Rank 2: 🎓 Higher Education / School Fees",
        "Rank 3: 💍 Marriage / Family Function",
        "Rank 4: 💼 Business Capital / Expansion",
        "Rank 5: 👤 Personal Savings / Investment"
    ], bg_color=LIGHT_BLUE, border_color=BORDER_BLUE)

    # =========================================================================
    # SLIDE 5: Objectives
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    add_header(s5, "Project Objectives & Core Goals", "4. CORE OBJECTIVES")
    add_card(s5, Inches(0.8), Inches(1.6), Inches(5.6), Inches(2.4), "🛡️ 1. Complete Security", [
        "256-bit SSL & JWT token authentication.",
        "Role-based access control (Member vs Admin).",
        "Mandatory 100% digital KYC verification."
    ], bg_color=LIGHT_BLUE, border_color=BORDER_BLUE, font_size=12)
    add_card(s5, Inches(6.8), Inches(1.6), Inches(5.6), Inches(2.4), "📊 2. Financial Transparency", [
        "Clear auction breakdown (Prize, Discount, Dividend).",
        "Month 1 Foreman & Month N Last Member guarantee.",
        "Real-time installment ledger tracking."
    ], bg_color=GREEN_BG, border_color=GREEN, font_size=12)
    add_card(s5, Inches(0.8), Inches(4.3), Inches(5.6), Inches(2.4), "⚡ 3. Automation & Convenience", [
        "Razorpay Instant UPI & QR Code checkout.",
        "Automated 10th monthly due email reminders.",
        "Instant digital receipt generation."
    ], bg_color=AMBER_BG, border_color=GOLD, font_size=12)
    add_card(s5, Inches(6.8), Inches(4.3), Inches(5.6), Inches(2.4), "⚖️ 4. Emergency Fairness", [
        "Prioritizes medical & educational needs.",
        "Removes bias in auction emergency bids.",
        "Complete audit trail of all transactions."
    ], bg_color=SLATE_BG, border_color=SLATE_DARK, font_size=12)

    # =========================================================================
    # SLIDE 6: System Modules
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    add_header(s6, "Core System Modules & Responsibilities", "5. SYSTEM MODULES")
    add_card(s6, Inches(0.8), Inches(1.6), Inches(3.6), Inches(2.4), "1. Auth & Security Module", [
        "User Registration & Login",
        "JWT Authentication Tokens",
        "Dual-Role Access Isolation"
    ], bg_color=LIGHT_BLUE, border_color=BORDER_BLUE, font_size=12)
    add_card(s6, Inches(4.8), Inches(1.6), Inches(3.6), Inches(2.4), "2. KYC & Guarantor Module", [
        "Aadhaar / PAN Document Upload",
        "Guarantor Verification Details",
        "Admin KYC Approval Portal"
    ], bg_color=GREEN_BG, border_color=GREEN, font_size=12)
    add_card(s6, Inches(8.8), Inches(1.6), Inches(3.6), Inches(2.4), "3. Chit Scheme Management", [
        "Scheme Creation (Bronze/Silver/Gold)",
        "Member Enrollment & Tickets",
        "Duration & Monthly Dues Config"
    ], bg_color=AMBER_BG, border_color=GOLD, font_size=12)
    add_card(s6, Inches(0.8), Inches(4.3), Inches(3.6), Inches(2.4), "4. Auction & Dividend Engine", [
        "Month 1 Foreman Company Pool",
        "Months 2..N-1 Auction Bidding",
        "5% Commission & Dividend Calc"
    ], bg_color=AMBER_BG, border_color=GOLD, font_size=12)
    add_card(s6, Inches(4.8), Inches(4.3), Inches(3.6), Inches(2.4), "5. Installment & Payment", [
        "Active Dues Identification",
        "Razorpay UPI / QR Checkout",
        "Verified Receipts (`REC-2026`)"
    ], bg_color=LIGHT_BLUE, border_color=BORDER_BLUE, font_size=12)
    add_card(s6, Inches(8.8), Inches(4.3), Inches(3.6), Inches(2.4), "6. Admin & Notifications", [
        "Payment Approval / Verification",
        "Automated 10th Due Reminders",
        "System Audit Logs & Reports"
    ], bg_color=GREEN_BG, border_color=GREEN, font_size=12)

    # =========================================================================
    # SLIDE 7: System Architecture
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    add_header(s7, "3-Tier System Architecture Specification", "6. SYSTEM ARCHITECTURE")
    add_card(s7, Inches(0.8), Inches(1.6), Inches(3.6), Inches(5.2), "Frontend Tier (Client)", [
        "React 18 Single Page App.",
        "Vite Build Tooling.",
        "TailwindCSS Styling.",
        "Axios REST HTTP Client.",
        "Lucide Icons & Confetti."
    ], bg_color=LIGHT_BLUE, border_color=BORDER_BLUE)
    add_card(s7, Inches(4.8), Inches(1.6), Inches(3.6), Inches(5.2), "Application Tier (API)", [
        "FastAPI Asynchronous Engine.",
        "Python 3.11 Runtime.",
        "SQLAlchemy ORM Layer.",
        "JWT Bearer Middleware.",
        "SMTP Email Service."
    ], bg_color=GREEN_BG, border_color=GREEN)
    add_card(s7, Inches(8.8), Inches(1.6), Inches(3.6), Inches(5.2), "Database Tier (Storage)", [
        "PostgreSQL 15 Relational DB.",
        "ACID Compliant Transactions.",
        "Encrypted Passwords (Bcrypt).",
        "Static Uploads Directory.",
        "Audit Log Storage."
    ], bg_color=SLATE_BG, border_color=SLATE_DARK)

    # =========================================================================
    # SLIDE 8: Use Case Diagram
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    add_header(s8, "Use Case Diagram & Actor Specifications", "7. USE CASE DESIGN")
    add_card(s8, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2), "Member Actor Use Cases", [
        "UC-1: Self Registration & Role-based Login",
        "UC-2: Upload Digital KYC Documents & Guarantor",
        "UC-3: Explore & Enroll in Auction Schemes",
        "UC-4: View Scheme Timeline & Net Payable Schedule",
        "UC-5: Pay Monthly Dues via Instant Razorpay UPI",
        "UC-6: View, Download & Print Verified Receipts"
    ], bg_color=LIGHT_BLUE, border_color=BORDER_BLUE)
    add_card(s8, Inches(6.8), Inches(1.6), Inches(5.6), Inches(5.2), "Admin Actor Use Cases", [
        "UC-7: Review & Approve Member KYC Verification",
        "UC-8: Create, Configure & Launch Chit Plans",
        "UC-9: Conduct Monthly Auction & Winner Selection",
        "UC-10: Review Emergency Priority Bids (Rank 1-5)",
        "UC-11: Verify & Approve Online Member Payments",
        "UC-12: Trigger 10th Monthly Due Email Reminders"
    ], bg_color=GREEN_BG, border_color=GREEN)

    # =========================================================================
    # SLIDE 9: DFD Diagram
    # =========================================================================
    s9 = prs.slides.add_slide(blank_layout)
    add_header(s9, "Data Flow Diagrams (DFD Level 0 & Level 1)", "8. PROCESS DATA FLOW")
    add_card(s9, Inches(0.8), Inches(1.6), Inches(11.6), Inches(2.4), "DFD Level 0 (Context Level Flow)", [
        "Member Entity ──> Submit KYC / Join Scheme / Pay Dues ──> SecureChit System",
        "System ──> Digital Receipts / Dividend Breakdown / Dues Reminders ──> Member Entity",
        "Admin Entity ──> Create Schemes / Approve KYC / Verify Payments ──> SecureChit System",
        "System ──> System Audit Logs / Financial Reports / Emergency Rankings ──> Admin Entity"
    ], bg_color=AMBER_BG, border_color=GOLD, font_size=12)
    add_card(s9, Inches(0.8), Inches(4.3), Inches(11.6), Inches(2.4), "DFD Level 1 (Subsystem Level Flow)", [
        "1.0 Auth Subsystem ──> Validates Credentials & Mints JWT Token ──> D1 (Users DB Store)",
        "2.0 Member & Scheme Subsystem ──> Stores Enrolled Members & Tickets ──> D2/D3 (Plans DB)",
        "3.0 Payment Gateway Subsystem ──> Processes UPI Payment & Generates Receipt ──> D5/D6 (Payments DB)",
        "4.0 Auction Engine Subsystem ──> Computes Bids & Calculates Dividends ──> D4 (MonthlyDraws DB)"
    ], bg_color=LIGHT_BLUE, border_color=BORDER_BLUE, font_size=12)

    # =========================================================================
    # SLIDE 10: Database / ER Diagram
    # =========================================================================
    s10 = prs.slides.add_slide(blank_layout)
    add_header(s10, "Database Entity-Relationship (ER) Schema", "9. DATABASE SPECIFICATION")
    add_card(s10, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2), "Database Tables Schema", [
        "1. `users`: User account, role (member/admin), kyc_status.",
        "2. `chit_plans`: Scheme config (monthly_amount, pool).",
        "3. `plan_members`: Member enrollments & ticket IDs.",
        "4. `monthly_draws`: Auction bids, dividend, commission.",
        "5. `installments`: Monthly dues (month_number, total_amount).",
        "6. `payments`: Transaction records (payment_method, tx_id)."
    ], bg_color=LIGHT_BLUE, border_color=BORDER_BLUE)
    add_card(s10, Inches(6.8), Inches(1.6), Inches(5.6), Inches(5.2), "ER Cardinality & Relationships", [
        "USERS (1) ────< (N) PLAN_MEMBERS (Members join schemes)",
        "CHIT_PLANS (1) ────< (N) PLAN_MEMBERS (Scheme has members)",
        "CHIT_PLANS (1) ────< (N) MONTHLY_DRAWS (Scheme has 1 draw/mo)",
        "PLAN_MEMBERS (1) ────< (N) INSTALLMENTS (Member has dues)",
        "INSTALLMENTS (1) ────< (N) PAYMENTS (Dues generate payments)",
        "USERS (1) ────< (N) PAYMENTS (User makes payments)"
    ], bg_color=GREEN_BG, border_color=GREEN)

    # =========================================================================
    # SLIDE 11: Key Features
    # =========================================================================
    s11 = prs.slides.add_slide(blank_layout)
    add_header(s11, "Key System Features & Innovations", "10. FEATURE HIGHLIGHTS")
    add_card(s11, Inches(0.8), Inches(1.6), Inches(3.6), Inches(5.2), "💳 Real-Time Payment Tracker", [
        "Active dues identification.",
        "Instant Razorpay UPI checkout.",
        "Google Pay / PhonePe QR scan.",
        "Immediate ledger updates."
    ], bg_color=GREEN_BG, border_color=GREEN)
    add_card(s11, Inches(4.8), Inches(1.6), Inches(3.6), Inches(5.2), "🧾 Verified Digital Receipts", [
        "Unique receipt numbers (`REC-2026`).",
        "Member name & ticket ID proof.",
        "Verified & Paid authorization badge.",
        "One-click browser print."
    ], bg_color=LIGHT_BLUE, border_color=BORDER_BLUE)
    add_card(s11, Inches(8.8), Inches(1.6), Inches(3.6), Inches(5.2), "📧 Automated Due Reminders", [
        "Triggered on the 10th monthly.",
        "Informs members before 15th draw.",
        "SMTP email integration."
    ], bg_color=AMBER_BG, border_color=GOLD)

    # =========================================================================
    # SLIDE 12: Technology Stack
    # =========================================================================
    s12 = prs.slides.add_slide(blank_layout)
    add_header(s12, "Technology Stack Specification", "11. TECH STACK")
    add_card(s12, Inches(0.8), Inches(1.6), Inches(3.6), Inches(5.2), "Frontend Architecture", [
        "React.js 18 & Vite",
        "React Router DOM v7",
        "TailwindCSS Styling",
        "Lucide React Icons",
        "Axios REST Client"
    ], bg_color=LIGHT_BLUE, border_color=BORDER_BLUE)
    add_card(s12, Inches(4.8), Inches(1.6), Inches(3.6), Inches(5.2), "Backend Architecture", [
        "Python 3.11 Runtime",
        "FastAPI Framework",
        "Uvicorn ASGI Server",
        "SQLAlchemy ORM Layer",
        "JWT & Bcrypt Security"
    ], bg_color=GREEN_BG, border_color=GREEN)
    add_card(s12, Inches(8.8), Inches(1.6), Inches(3.6), Inches(5.2), "Database & Testing Tooling", [
        "PostgreSQL 15 Database",
        "Alembic DB Migrations",
        "PyTest Automated Suite",
        "Puppeteer E2E Testing",
        "Node.js Integration Suite"
    ], bg_color=SLATE_BG, border_color=SLATE_DARK)

    # =========================================================================
    # SLIDE 13: Testing
    # =========================================================================
    s13 = prs.slides.add_slide(blank_layout)
    add_header(s13, "System Testing & Quality Assurance", "12. TESTING & VALIDATION")
    add_card(s13, Inches(0.8), Inches(1.6), Inches(11.6), Inches(5.2), "Mandatory 16-Point Integration Test Verification Suite (100% Passed)", [
        "Group 1: Member Authentication & Dues Identification ──> [PASS] 4/4 Tests",
        "Group 2: Payment Gateway & Request Submission ──> [PASS] 4/4 Tests",
        "Group 3: Admin Verification & Approval Lifecycle ──> [PASS] 4/4 Tests",
        "Group 4: Integration Verification & Digital Receipt Generation ──> [PASS] 4/4 Tests",
        "",
        "Test Execution Result: 16/16 Tests Passed (0 Failed) | Executed via `node test_payment_installment_integration.cjs`"
    ], bg_color=GREEN_BG, border_color=GREEN, font_size=13)

    # =========================================================================
    # SLIDE 14: Results / Screenshots
    # =========================================================================
    s14 = prs.slides.add_slide(blank_layout)
    add_header(s14, "System Results & Interface Screenshots", "13. SYSTEM RESULTS")
    add_card(s14, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2), "Member Installment & Gateway UI", [
        "Displays active dues & pending amount.",
        "Scannable Google Pay UPI QR code modal.",
        "Emergency reason rank selection dropdown.",
        "Instant request submission to Admin."
    ], bg_color=LIGHT_BLUE, border_color=BORDER_BLUE)
    add_card(s14, Inches(6.8), Inches(1.6), Inches(5.6), Inches(5.2), "Verified Payment Receipt Modal", [
        "Displays receipt number (`REC-2026-0014`).",
        "Shows verified member name & scheme details.",
        "Displays paid amount ₹2,000 & payment ref ID.",
        "Includes printable official PDF receipt voucher."
    ], bg_color=GREEN_BG, border_color=GREEN)

    # =========================================================================
    # SLIDE 15: Future Enhancement
    # =========================================================================
    s15 = prs.slides.add_slide(blank_layout)
    add_header(s15, "Future System Enhancements", "14. FUTURE ROADMAP")
    add_card(s15, Inches(0.8), Inches(1.6), Inches(5.6), Inches(2.4), "📱 Mobile Application", [
        "Cross-platform React Native / Flutter app.",
        "Push notifications for due dates.",
        "Biometric fingerprint / FaceID login."
    ], bg_color=LIGHT_BLUE, border_color=BORDER_BLUE, font_size=12)
    add_card(s15, Inches(6.8), Inches(1.6), Inches(5.6), Inches(2.4), "⚡ Live Bidding WebSockets", [
        "Real-time interactive auction room.",
        "Live auction bid ticker & timer.",
        "Auto-bidding proxy for members."
    ], bg_color=GREEN_BG, border_color=GREEN, font_size=12)
    add_card(s15, Inches(0.8), Inches(4.3), Inches(5.6), Inches(2.4), "🔗 Decentralized Blockchain Ledger", [
        "Immutable transaction record on Polygon.",
        "Smart contract automated payouts.",
        "100% decentralized trust verification."
    ], bg_color=AMBER_BG, border_color=GOLD, font_size=12)
    add_card(s15, Inches(6.8), Inches(4.3), Inches(5.6), Inches(2.4), "💬 WhatsApp Business Bot", [
        "Instant WhatsApp payment receipts.",
        "Automated due reminders via WhatsApp API.",
        "One-click UPI payment links."
    ], bg_color=SLATE_BG, border_color=SLATE_DARK, font_size=12)

    # =========================================================================
    # SLIDE 16: Conclusion
    # =========================================================================
    s16 = prs.slides.add_slide(blank_layout)
    add_header(s16, "Project Conclusion", "15. CONCLUSION")
    add_card(s16, Inches(0.8), Inches(1.6), Inches(11.6), Inches(5.2), "Summary of Accomplishments", [
        "Successfully developed and deployed the Secure Chit Fund Management System (SecureChit).",
        "Achieved 100% seamless integration between Payment Gateway and Installment Tracker Modules.",
        "Eliminated traditional manual book-keeping vulnerabilities with automated 5% commission & dividend schedules.",
        "Implemented robust 3-tier architecture with FastAPI, React.js, and PostgreSQL.",
        "Validated with a 16-point automated integration test verification suite (16/16 Passed)."
    ], bg_color=GREEN_BG, border_color=GREEN, font_size=13)

    # =========================================================================
    # SLIDE 17: Thank You
    # =========================================================================
    s17 = prs.slides.add_slide(blank_layout)
    bg17 = s17.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg17.fill.solid()
    bg17.fill.fore_color.rgb = NAVY
    bg17.line.color.rgb = NAVY

    tb17 = s17.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(11.333), Inches(3.5))
    tf17 = tb17.text_frame
    tf17.word_wrap = True

    p0 = tf17.paragraphs[0]
    p0.text = "THANK YOU!"
    p0.alignment = PP_ALIGN.CENTER
    p0.font.size = Pt(48)
    p0.font.bold = True
    p0.font.color.rgb = GOLD

    p1 = tf17.add_paragraph()
    p1.text = "\nQuestions & Discussion"
    p1.alignment = PP_ALIGN.CENTER
    p1.font.size = Pt(24)
    p1.font.bold = True
    p1.font.color.rgb = WHITE

    p2 = tf17.add_paragraph()
    p2.text = "\nSecure Chit Fund Management System"
    p2.alignment = PP_ALIGN.CENTER
    p2.font.size = Pt(16)
    p2.font.color.rgb = WHITE

    # Try saving to primary file or updated file if locked
    output_path = r"c:\chit fund\SecureChit_Overall_Project_Presentation_Updated.pptx"
    prs.save(output_path)
    print(f"[SUCCESS] Updated presentation saved to {output_path}")

    try:
        primary_path = r"c:\chit fund\SecureChit_Overall_Project_Presentation.pptx"
        prs.save(primary_path)
        print(f"[SUCCESS] Also updated {primary_path}")
    except Exception as e:
        print(f"[NOTE] Primary file locked by PowerPoint viewer, saved clean file as {output_path}")

if __name__ == "__main__":
    build_overall_presentation()
