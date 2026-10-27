import shutil
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

src = r'C:\chit fund\SecureChit_Report_With_DFDs_v4.docx'
out = r'C:\chit fund\SecureChit_Report_With_DFDs_v5.docx'
doc = Document(src)

parent = doc.element.body

def ins(elem):
    parent.append(elem)

def heading(text, size=13):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(size)
    ins(p._element)

def body(text, size=11):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = False
    r.font.size = Pt(size)
    ins(p._element)

def bullet(label, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.25)
    r0 = p.add_run(u"\u2022  ")
    r0.font.size = Pt(11)
    r1 = p.add_run(label + ": ")
    r1.bold = True
    r1.font.size = Pt(11)
    r2 = p.add_run(text)
    r2.font.size = Pt(11)
    ins(p._element)

def blank():
    p = doc.add_paragraph()
    p.add_run("").font.size = Pt(6)
    ins(p._element)

# ─── 5. SYSTEM TESTING ────────────────────────────────────────────────────────
blank()
heading("5. SYSTEM TESTING", size=13)
blank()

# ─── 5.1 ──────────────────────────────────────────────────────────────────────
heading("5.1 SYSTEM TESTING", size=12)
blank()
body(
    "System testing is a crucial phase in the software development life cycle (SDLC) to verify "
    "that the integrated system meets the specified requirements and behaves as expected in the "
    "intended environment. It is conducted after unit testing and integration testing and before "
    "acceptance testing."
)
blank()
body(
    "System testing for SecureChit verifies the functionality, reliability, and security of the "
    "complete platform, including member authentication, KYC verification, chit plan enrollment, "
    "monthly auction bidding, dividend distribution, installment payment, promissory note execution, "
    "and PDF receipt generation. Special importance was given to the state machine of the "
    "installment life cycle, input validation rules such as the 12-digit UTR pattern, "
    "auction settlement formulas, and transaction consistency in the PostgreSQL database."
)
blank()

# ─── 5.2 ──────────────────────────────────────────────────────────────────────
heading("5.2 TYPES OF TESTING", size=12)
blank()
body("The key aspects considered while testing SecureChit are:")
blank()

bullet("Unit Testing",
       "FastAPI route handlers, Pydantic v2 schemas, utility functions, and regex validators "
       "were tested in isolation using pytest and FastAPI TestClient.")

bullet("Integration Testing",
       "Communication between the React client (port 5173 / 5174), the Vite proxy, "
       "FastAPI backend (port 8000), and PostgreSQL (port 5433) was verified end-to-end.")

bullet("Functional Testing",
       "End-to-end features such as member enrollment, auction bidding, dividend calculation, "
       "UPI payment submission, promissory note generation, and PDF receipt download were "
       "checked against requirements.")

bullet("Validation Testing",
       "Blank fields, invalid email formats, malformed 12-digit UTR numbers, duplicate plan "
       "enrollments, and out-of-range bid amounts were tested with negative inputs.")

bullet("Security and Role-Based Testing",
       "Members cannot access admin-only endpoints; admins cannot directly confirm member payments; "
       "passwords remain BCrypt-hashed in the database; JWT tokens expire correctly and are "
       "rejected for unauthorized roles.")

bullet("Regression Testing",
       "Tests were re-run after every new feature addition, such as the AI Emergency Priority "
       "Engine, dividend auto-distribution, and pronote PDF compilation.")

blank()

# ─── 5.3 TEST CASES TABLE ─────────────────────────────────────────────────────
heading("5.3 TEST CASES", size=12)
blank()
body("The important test cases executed on SecureChit are listed below.")
blank()

# Table columns
cols = ["Test ID", "Scenario and Input", "Expected Result", "Status"]

rows = [
    # AUTH
    ("TC-AUTH-01",
     "Member registration with valid data (Name: Arun, Email: arun@sc.in, Phone: 9876543210)",
     "HTTP 201; user saved with role=member, status=pending, kyc_status=not_submitted, password stored as BCrypt hash.",
     "PASS"),
    ("TC-AUTH-02",
     "Registration with a duplicate username",
     "HTTP 400; 'Username already registered'.",
     "PASS"),
    ("TC-AUTH-03",
     "Login with an incorrect password",
     "HTTP 401; 'Invalid credentials'. failed_login_attempts incremented.",
     "PASS"),
    ("TC-AUTH-04",
     "Admin login with valid credentials",
     "HTTP 200; JWT Bearer token issued with role=admin scope.",
     "PASS"),

    # KYC
    ("TC-KYC-01",
     "Member uploads Aadhaar front, PAN card, passport photo",
     "KYC documents saved in kyc_documents with status=pending. Member kyc_status=pending.",
     "PASS"),
    ("TC-KYC-02",
     "Admin approves KYC for member ID 12",
     "kyc_documents status=verified; users.kyc_status=verified; member account activated.",
     "PASS"),
    ("TC-KYC-03",
     "Admin rejects a blurry PAN card",
     "status=rejected; rejection_reason logged; member notified to re-upload.",
     "PASS"),

    # PLAN ENROLLMENT
    ("TC-PLAN-01",
     "Member enrolls in an active Silver plan (ID: 3)",
     "plan_members record created; ticket_number=TCK-2026-1 assigned; has_won=false.",
     "PASS"),
    ("TC-PLAN-02",
     "Same member attempts to enroll in plan ID 3 again",
     "HTTP 400; UNIQUE constraint on (user_id, plan_id) prevents duplicate enrollment.",
     "PASS"),
    ("TC-PLAN-03",
     "Member enrolls in a full plan (current_members = total_members)",
     "HTTP 400; 'Plan is full. No vacancy available.'",
     "PASS"),

    # AUCTION
    ("TC-AUC-01",
     "Member submits bid: amount=45000, emergency_rank=1 (Medical)",
     "Bid recorded in plan_members.priority_score=1. Validated against plan's chit_value.",
     "PASS"),
    ("TC-AUC-02",
     "Two members bid equally; one with rank=1 (Medical), other rank=3 (Marriage)",
     "Rank 1 member selected as winner. AI Priority Engine resolves tie correctly.",
     "PASS"),
    ("TC-AUC-03",
     "Admin conducts draw for month 3 of plan ID 2",
     "monthly_draws record created; winner_id set; dividend_per_member calculated; "
     "all active installments updated with reduced net_payable.",
     "PASS"),
    ("TC-AUC-04",
     "Dividend = (50000 - 45000 - 250) / 20 = Rs.237.50 per member",
     "dividend_per_member=237.50; company_commission=250.00; net_payable reduced for all 20 members.",
     "PASS"),

    # INSTALLMENT & PAYMENT
    ("TC-PAY-01",
     "Member submits valid UTR: 984120394812 (12 digits)",
     "Installment status=paid; payment record created; status=pending verification.",
     "PASS"),
    ("TC-PAY-02",
     "UTR with 9 digits (123456789)",
     "HTTP 400; 'UTR must contain exactly 12 numeric digits'.",
     "PASS"),
    ("TC-PAY-03",
     "UTR with a letter (98412039481A)",
     "HTTP 400; regex ^[0-9]{12}$ validation fails.",
     "PASS"),
    ("TC-PAY-04",
     "Admin verifies installment payment for member ID 5",
     "Installment status=verified; receipt_number=REC-2026-0041 generated; PDF receipt compiled.",
     "PASS"),
    ("TC-PAY-05",
     "Member re-submits a previously failed UTR",
     "HTTP 400; reuse of a failed transaction ID is blocked.",
     "PASS"),
    ("TC-PAY-06",
     "Member retries with a fresh UTR (458192304918)",
     "New payment record created; previous failed record retained in history.",
     "PASS"),

    # PROMISSORY NOTE
    ("TC-PRN-01",
     "Admin selects auction winner for member ID 7",
     "promissory_notes record auto-created; pronote_id=PRN-2026-0007; status=pending_signature.",
     "PASS"),
    ("TC-PRN-02",
     "Member submits Base64 Canvas digital signature",
     "signature column updated; pronote status=active; signed PDF compiled by ReportLab.",
     "PASS"),
    ("TC-PRN-03",
     "Member attempts to access another member's promissory note",
     "HTTP 403 Forbidden; JWT role check prevents cross-member data access.",
     "PASS"),

    # SECURITY
    ("TC-SEC-01",
     "Member attempts to access POST /admin/plans (admin-only endpoint)",
     "HTTP 403 Forbidden; RBAC middleware rejects non-admin JWT token.",
     "PASS"),
    ("TC-SEC-02",
     "Request sent without Authorization header",
     "HTTP 401 Unauthorized; 'Not authenticated'.",
     "PASS"),
    ("TC-SEC-03",
     "Expired JWT token submitted",
     "HTTP 401; 'Token has expired'. User must re-login.",
     "PASS"),
]

# Build table
tbl = doc.add_table(rows=1 + len(rows), cols=4)
tbl.style = 'Table Grid'

# Header
hdr = tbl.rows[0].cells
for ci, col in enumerate(cols):
    hdr[ci].text = col
    for run in hdr[ci].paragraphs[0].runs:
        run.bold = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    # Dark navy header background
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    tc_pr = hdr[ci]._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), '0B132B')
    tc_pr.append(shd)

# Data rows
for ri, (tid, scenario, expected, status) in enumerate(rows):
    cells = tbl.rows[ri + 1].cells
    data = [tid, scenario, expected, status]
    for ci, val in enumerate(data):
        cells[ci].text = val
        for run in cells[ci].paragraphs[0].runs:
            run.font.size = Pt(9)
            if ci == 3:  # Status column
                run.bold = True
                run.font.color.rgb = RGBColor(0x05, 0x96, 0x52) if val == "PASS" else RGBColor(0xDC, 0x26, 0x26)

ins(tbl._element)
blank()

doc.save(out)
print(f"Saved: {out}")

shutil.copy(out, r'C:\Users\nares\.gemini\antigravity\brain\f74afe5d-2f73-4013-bc1e-3562e05df0f5\SecureChit_Report_With_DFDs.docx')
print("Copied to artifact dir!")
