import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_pptx(output_path):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette
    NAVY = RGBColor(26, 43, 73)          # #1A2B49
    GOLD = RGBColor(217, 119, 6)         # Amber-600
    SLATE_DARK = RGBColor(15, 23, 42)    # Slate-900
    LIGHT_BLUE = RGBColor(239, 246, 255) # Blue-50
    BORDER_BLUE = RGBColor(59, 130, 246) # Blue-500
    WHITE = RGBColor(255, 255, 255)
    GREEN = RGBColor(16, 185, 129)       # Emerald-500
    GREEN_BG = RGBColor(236, 253, 245)    # Emerald-50
    AMBER_BG = RGBColor(254, 243, 199)   # Amber-100
    GRAY_TEXT = RGBColor(71, 85, 105)

    def add_header(slide, title_text, category_text="SECURECHIT PAYMENT MODULE GUIDE"):
        tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(1.1))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p0 = tf.paragraphs[0]
        p0.text = category_text.upper()
        p0.font.size = Pt(11)
        p0.font.bold = True
        p0.font.color.rgb = GOLD
        
        p1 = tf.add_paragraph()
        p1.text = title_text
        p1.font.size = Pt(24)
        p1.font.bold = True
        p1.font.color.rgb = NAVY

    def add_card(slide, left, top, width, height, title, items, bg_color=LIGHT_BLUE, border_color=BORDER_BLUE, title_color=NAVY):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1.5)

        tf = shape.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.2)
        tf.margin_right = Inches(0.2)
        tf.margin_top = Inches(0.2)
        tf.margin_bottom = Inches(0.2)

        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = title_color

        for item in items:
            p = tf.add_paragraph()
            p.text = "• " + item
            p.font.size = Pt(12)
            p.font.color.rgb = SLATE_DARK
            p.space_before = Pt(4)

    # ----------------------------------------------------
    # SLIDE 1: Title Slide
    # ----------------------------------------------------
    slide1 = prs.slides.add_slide(blank_layout)
    bg1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = NAVY
    bg1.line.color.rgb = NAVY

    tb1 = slide1.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(11.333), Inches(3.5))
    tf1 = tb1.text_frame
    tf1.word_wrap = True

    p = tf1.paragraphs[0]
    p.text = "CHIT FUND PROJECT"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = GOLD

    p = tf1.add_paragraph()
    p.text = "Payment Module Implementation Notes"
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.space_before = Pt(10)

    p = tf1.add_paragraph()
    p.text = "End-to-End Guide: Database Schema, FastAPI APIs, React Frontend & Admin Verification Workflow"
    p.font.size = Pt(16)
    p.font.color.rgb = RGBColor(191, 219, 254)
    p.space_before = Pt(15)

    # ----------------------------------------------------
    # SLIDE 2: System Workflow & Architecture
    # ----------------------------------------------------
    slide2 = prs.slides.add_slide(blank_layout)
    add_header(slide2, "1. Payment Module Workflow Architecture")

    add_card(slide2, Inches(0.8), Inches(1.6), Inches(3.6), Inches(5.2), 
             "1. Member Actions", 
             ["View Monthly Dues via GET /dues",
              "Click 'Pay via Instant UPI / QR'",
              "Select Emergency Reason (Rank 1-5)",
              "Scan GPay / PhonePe QR Code",
              "Submit Payment Request"], 
             LIGHT_BLUE, BORDER_BLUE, NAVY)

    add_card(slide2, Inches(4.8), Inches(1.6), Inches(3.6), Inches(5.2), 
             "2. System & Database", 
             ["Payment created with status 'pending'",
              "Installment updated to 'paid'",
              "Receipt number generated",
              "Automated notification sent to Admin",
              "Frontend polls for verification"], 
             AMBER_BG, GOLD, NAVY)

    add_card(slide2, Inches(8.8), Inches(1.6), Inches(3.6), Inches(5.2), 
             "3. Admin Verification", 
             ["Admin checks payment details & UPI ID",
              "Approve (Status -> 'success')",
              "Installment status -> 'verified'",
              "Verified Receipt issued to Member",
              "Confetti animation on Frontend 🎉"], 
             GREEN_BG, GREEN, NAVY)

    # ----------------------------------------------------
    # SLIDE 3: Database Schema Design
    # ----------------------------------------------------
    slide3 = prs.slides.add_slide(blank_layout)
    add_header(slide3, "2. Database Schema (SQLAlchemy Models)")

    add_card(slide3, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2), 
             "Payments Table (`payments`)", 
             ["id: Primary Key (Integer)",
              "installment_id: Foreign Key (installments.id)",
              "user_id: Foreign Key (users.id)",
              "amount: Float (Monthly due amount)",
              "payment_method: String (e.g. Google Pay UPI)",
              "transaction_id: String (pay_upi_...)",
              "status: Enum ('pending', 'success', 'failed')",
              "notes: Text (Emergency Reason & notes)",
              "created_at: DateTime Timestamp"], 
             LIGHT_BLUE, BORDER_BLUE, NAVY)

    add_card(slide3, Inches(6.8), Inches(1.6), Inches(5.6), Inches(5.2), 
             "Installments Table Fields (`installments`)", 
             ["id: Primary Key (Integer)",
              "status: Enum ('pending', 'paid', 'verified', 'overdue')",
              "receipt_number: String (e.g. REC-2026-0001)",
              "verified_by: Foreign Key (users.id)",
              "verified_at: DateTime Timestamp",
              "paid_date: Date",
              "notes / transaction details link"], 
             AMBER_BG, GOLD, NAVY)

    # ----------------------------------------------------
    # SLIDE 4: FastAPI Backend APIs
    # ----------------------------------------------------
    slide4 = prs.slides.add_slide(blank_layout)
    add_header(slide4, "3. Backend API Endpoints (FastAPI)")

    add_card(slide4, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2), 
             "Member API Router (`/member/payments`)", 
             ["GET /dues: List active unpaid dues",
              "POST /razorpay/create-order: Generate Razorpay Order ID & metadata",
              "POST /razorpay/verify: Submit payment request & emergency reason for verification",
              "GET /history: Fetch member's past payment receipts & status"], 
             LIGHT_BLUE, BORDER_BLUE, NAVY)

    add_card(slide4, Inches(6.8), Inches(1.6), Inches(5.6), Inches(5.2), 
             "Admin API Router (`/admin/payments`)", 
             ["GET /: List all member payments (filterable by pending/success/failed)",
              "PUT /{id}/status: Approve ('success' -> verified) or Reject ('failed' -> pending)",
              "POST /send-due-reminders: Send monthly 10th due reminders",
              "DELETE /clear-all: Reset payments for test mode"], 
             GREEN_BG, GREEN, NAVY)

    # ----------------------------------------------------
    # SLIDE 5: Frontend React Implementation
    # ----------------------------------------------------
    slide5 = prs.slides.add_slide(blank_layout)
    add_header(slide5, "4. React Frontend Components & UI")

    add_card(slide5, Inches(0.8), Inches(1.6), Inches(3.6), Inches(5.2), 
             "Payments Page", 
             ["File: `Payments.jsx`",
              "Displays member dues list",
              "Shows transaction history",
              "Status badges (Verified, Pending, Rejected)",
              "Triggers Pay modal"], 
             LIGHT_BLUE, BORDER_BLUE, NAVY)

    add_card(slide5, Inches(4.8), Inches(1.6), Inches(3.6), Inches(5.2), 
             "Razorpay / UPI Modal", 
             ["File: `RazorpayModal.jsx`",
              "GPay QR Code & UPI Copy",
              "Emergency Reason selector",
              "Submit request button",
              "Background status polling"], 
             AMBER_BG, GOLD, NAVY)

    add_card(slide5, Inches(8.8), Inches(1.6), Inches(3.6), Inches(5.2), 
             "Digital Receipt Modal", 
             ["File: `PaymentReceiptModal.jsx`",
              "Formal digital receipt format",
              "Receipt # (REC-2026-XXXX)",
              "VERIFIED stamp & details",
              "Print / Save PDF option"], 
             GREEN_BG, GREEN, NAVY)

    # ----------------------------------------------------
    # SLIDE 6: Summary & Best Practices
    # ----------------------------------------------------
    slide6 = prs.slides.add_slide(blank_layout)
    add_header(slide6, "5. Best Practices & Key Takeaways for Friend's Project")

    add_card(slide6, Inches(0.8), Inches(1.6), Inches(11.7), Inches(5.2), 
             "Core Guidelines & Architecture Highlights", 
             ["2-Step Verification Model: Prevents fraudulent claims by requiring Admin verification before final receipt generation.",
              "Emergency Reason Priority Score: Rank 1 (Hospital) to Rank 5 (Personal) helps prioritize members during chit bidding/draws.",
              "Polling Mechanism: Frontend automatically polls backend history every 3s to detect Admin approvals seamlessly.",
              "Role-Based Access Control: Admin endpoints guarded with `require_admin` dependency for maximum security.",
              "UX Delights: Confetti animation (`canvas-confetti`) on payment approval enhances user experience."], 
             LIGHT_BLUE, BORDER_BLUE, NAVY)

    prs.save(output_path)
    print(f"Successfully generated PPTX: {output_path}")

def create_pdf(output_path):
    from reportlab.lib.pagesizes import letter
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib import colors

    doc = SimpleDocTemplate(
        output_path,
        pagesize=letter,
        rightMargin=40,
        leftMargin=40,
        topMargin=40,
        bottomMargin=40
    )

    styles = getSampleStyleSheet()
    
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=22,
        leading=26,
        textColor=colors.HexColor('#1A2B49'),
        spaceAfter=6
    )

    subtitle_style = ParagraphStyle(
        'DocSubTitle',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=12,
        leading=16,
        textColor=colors.HexColor('#D97706'),
        spaceAfter=15
    )

    heading_style = ParagraphStyle(
        'SectionHeading',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=15,
        leading=18,
        textColor=colors.HexColor('#1A2B49'),
        spaceBefore=14,
        spaceAfter=8
    )

    body_style = ParagraphStyle(
        'BodyDark',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=10,
        leading=14,
        textColor=colors.HexColor('#1E293B'),
        spaceAfter=6
    )

    bullet_style = ParagraphStyle(
        'BulletText',
        parent=body_style,
        leftIndent=15,
        spaceAfter=4
    )

    story = []

    story.append(Paragraph("Chit Fund Project - Payment Module Notes", title_style))
    story.append(Paragraph("Complete Architecture, Database Schema, FastAPI APIs & React Frontend Guide", subtitle_style))
    story.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor('#3B82F6'), spaceAfter=15))

    # Section 1
    story.append(Paragraph("1. System Architecture & Workflow", heading_style))
    story.append(Paragraph("The Payment Module operates on a <b>2-Step Verification Model</b> to ensure secure transactions and valid payment receipts:", body_style))
    story.append(Paragraph("• <b>Step 1: Member Views Dues</b> - Member checks active monthly dues via <code>GET /api/v1/member/payments/dues</code>.", bullet_style))
    story.append(Paragraph("• <b>Step 2: Emergency Reason & QR Scan</b> - Member clicks <i>Pay via Instant UPI / QR</i>, selects emergency reason (Rank 1-5), and scans GPay QR Code.", bullet_style))
    story.append(Paragraph("• <b>Step 3: Request Submission</b> - Calls <code>POST /razorpay/verify</code>. Payment record created with status <b>pending</b>, installment status updated to <b>paid</b>.", bullet_style))
    story.append(Paragraph("• <b>Step 4: Admin Verification</b> - Admin reviews transaction ID and UPI ID in Admin Panel. Clicking Approve sets payment status to <b>success</b> and installment to <b>verified</b>.", bullet_style))
    story.append(Paragraph("• <b>Step 5: Verified Receipt</b> - Formal receipt (REC-2026-XXXX) is issued to member with confetti celebration.", bullet_style))

    # Section 2
    story.append(Paragraph("2. Database Schema (SQLAlchemy)", heading_style))
    
    table_data = [
        [Paragraph("<b>Field Name</b>", body_style), Paragraph("<b>Data Type</b>", body_style), Paragraph("<b>Description</b>", body_style)],
        [Paragraph("id", body_style), Paragraph("Integer (PK)", body_style), Paragraph("Unique payment identifier", body_style)],
        [Paragraph("installment_id", body_style), Paragraph("Integer (FK)", body_style), Paragraph("References installments.id", body_style)],
        [Paragraph("user_id", body_style), Paragraph("Integer (FK)", body_style), Paragraph("References users.id", body_style)],
        [Paragraph("amount", body_style), Paragraph("Float", body_style), Paragraph("Monthly installment payment amount", body_style)],
        [Paragraph("payment_method", body_style), Paragraph("String(50)", body_style), Paragraph("e.g. Google Pay UPI (nareshk1213@oksbi)", body_style)],
        [Paragraph("transaction_id", body_style), Paragraph("String(100)", body_style), Paragraph("Unique UPI payment / transaction ID", body_style)],
        [Paragraph("status", body_style), Paragraph("Enum", body_style), Paragraph("pending | success | failed", body_style)],
        [Paragraph("notes", body_style), Paragraph("Text", body_style), Paragraph("Emergency reason & Admin notes", body_style)],
    ]

    t = Table(table_data, colWidths=[110, 90, 310])
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#EFF6FF')),
        ('BOX', (0, 0), (-1, -1), 1, colors.HexColor('#3B82F6')),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#CBD5E1')),
        ('TOPPADDING', (0, 0), (-1, -1), 5),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 5),
    ]))
    story.append(t)

    # Section 3
    story.append(Paragraph("3. FastAPI Backend Endpoints", heading_style))
    story.append(Paragraph("<b>Member Router (<code>/api/v1/member/payments</code>):</b>", body_style))
    story.append(Paragraph("• <code>GET /dues</code> - List all unpaid installments for current member.", bullet_style))
    story.append(Paragraph("• <code>POST /razorpay/create-order</code> - Generate Razorpay Order ID for test mode checkout.", bullet_style))
    story.append(Paragraph("• <code>POST /razorpay/verify</code> - Submit member payment request for Admin verification.", bullet_style))
    story.append(Paragraph("• <code>GET /history</code> - Retrieve member payment receipt history.", bullet_style))

    story.append(Paragraph("<b>Admin Router (<code>/api/v1/admin/payments</code>):</b>", body_style))
    story.append(Paragraph("• <code>GET /</code> - List all payments with filter by status (pending/success/failed).", bullet_style))
    story.append(Paragraph("• <code>PUT /{id}/status</code> - Update status (success -> marks installment verified).", bullet_style))
    story.append(Paragraph("• <code>POST /send-due-reminders</code> - Trigger monthly due notification reminders.", bullet_style))

    # Section 4
    story.append(Paragraph("4. React Frontend Components", heading_style))
    story.append(Paragraph("• <b><code>Payments.jsx</code></b>: Displays active dues & past transaction history with status badges.", bullet_style))
    story.append(Paragraph("• <b><code>RazorpayModal.jsx</code></b>: Modal featuring GPay QR code, UPI copy, emergency reason dropdown, and polling.", bullet_style))
    story.append(Paragraph("• <b><code>PaymentReceiptModal.jsx</code></b>: Modal displaying official verified digital receipt with print capability.", bullet_style))

    doc.build(story)
    print(f"Successfully generated PDF: {output_path}")

if __name__ == "__main__":
    pptx_path = os.path.join(os.getcwd(), "ChitFund_Payment_Module_Guide.pptx")
    pdf_path = os.path.join(os.getcwd(), "ChitFund_Payment_Module_Guide.pdf")

    create_pptx(pptx_path)
    create_pdf(pdf_path)
