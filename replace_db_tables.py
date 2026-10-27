import shutil
from docx import Document
from docx.shared import Pt, Inches, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

src = r'C:\chit fund\SecureChit_Report_With_DFDs_v7.docx'
out = r'C:\chit fund\SecureChit_Report_With_DFDs_v8.docx'
shutil.copy(src, out)
doc = Document(out)

# ── helpers ──────────────────────────────────────────────────────────────────
def set_cell_bg(cell, hex_color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tc_pr.append(shd)

def set_row_height(row, height_cm):
    tr = row._tr
    trPr = tr.get_or_add_trPr()
    trHeight = OxmlElement('w:trHeight')
    trHeight.set(qn('w:val'), str(int(height_cm * 567)))  # 1cm = 567 twips
    trHeight.set(qn('w:hRule'), 'atLeast')
    trPr.append(trHeight)

def make_table(doc, heading_text, rows_data):
    """Create heading para + 3-col gray-header table matching the photo style."""
    # Heading: bold italic "Table N: name"
    hp = doc.add_paragraph()
    hr = hp.add_run(heading_text)
    hr.bold = True
    hr.italic = True
    hr.font.size = Pt(11)

    # Table: 3 columns — Field | Type | Description
    tbl = doc.add_table(rows=1 + len(rows_data), cols=3)
    tbl.style = 'Table Grid'

    # Column widths (approximate A4 portrait usable width ~14cm)
    widths = [Cm(3.5), Cm(4.5), Cm(6.2)]

    # Header row
    hdr = tbl.rows[0].cells
    for ci, col_name in enumerate(["Field", "Type", "Description"]):
        hdr[ci].width = widths[ci]
        set_cell_bg(hdr[ci], 'BFBFBF')
        set_row_height(tbl.rows[0], 1.0)
        p = hdr[ci].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(col_name)
        r.bold = True
        r.font.size = Pt(10)

    # Data rows
    for ri, (field, ftype, desc) in enumerate(rows_data):
        row = tbl.rows[ri + 1]
        set_row_height(row, 0.85)
        cells = row.cells
        data = [field, ftype, desc]
        for ci, val in enumerate(data):
            cells[ci].width = widths[ci]
            p = cells[ci].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(val)
            r.font.size = Pt(10)

    return hp, tbl

# ── All 15 tables data (Field | Type with constraints merged | Description) ──

ALL_TABLES = [
    ("Table 1: users", [
        ("id",                    "INT, PK, Auto-Increment",      "Unique user identifier"),
        ("username",              "VARCHAR(50), UNIQUE, NOT NULL", "Login username"),
        ("name",                  "VARCHAR(100), NOT NULL",        "Full name"),
        ("email",                 "VARCHAR(150), NOT NULL",        "Email address"),
        ("phone",                 "VARCHAR(15), NOT NULL",         "Contact number"),
        ("password_hash",         "VARCHAR(255), NOT NULL",        "BCrypt hashed password"),
        ("role",                  "ENUM, NOT NULL",                "member / admin"),
        ("status",                "ENUM, NOT NULL",                "pending / approved / active / suspended / rejected"),
        ("aadhar_number",         "VARCHAR(20), NULL",             "Aadhaar card number"),
        ("pan_number",            "VARCHAR(20), NULL",             "PAN card number"),
        ("date_of_birth",         "DATE, NULL",                    "Date of birth"),
        ("occupation",            "VARCHAR(100), NULL",            "Member occupation"),
        ("kyc_status",            "ENUM, NOT NULL",                "not_submitted / pending / verified / rejected"),
        ("failed_login_attempts", "INT, DEFAULT 0",                "Brute-force login counter"),
        ("last_login",            "TIMESTAMP, NULL",               "Last successful login timestamp"),
        ("is_deleted",            "BOOLEAN, DEFAULT FALSE",        "Soft delete flag"),
        ("created_at",            "TIMESTAMP, Server Default",     "Registration timestamp"),
        ("updated_at",            "TIMESTAMP, Auto-update",        "Last modification timestamp"),
    ]),
    ("Table 2: kyc_documents", [
        ("id",               "INT, PK",                    "Unique KYC record ID"),
        ("user_id",          "INT, FK -> users.id",        "Document owner"),
        ("document_type",    "ENUM, NOT NULL",             "aadhar_front / aadhar_back / pan_card / address_proof / passport_photo"),
        ("file_path",        "VARCHAR(255), NOT NULL",     "Uploaded file storage path"),
        ("status",           "ENUM, DEFAULT pending",      "pending / verified / rejected"),
        ("rejection_reason", "TEXT, NULL",                 "Reason if rejected by admin"),
        ("verified_by",      "INT, FK -> users.id",        "Admin who verified the document"),
        ("verified_at",      "TIMESTAMP, NULL",            "Verification timestamp"),
        ("created_at",       "TIMESTAMP, Server Default",  "Upload timestamp"),
    ]),
    ("Table 3: chit_plans", [
        ("id",               "INT, PK",                   "Unique plan ID"),
        ("plan_name",        "VARCHAR(100), NOT NULL",     "Plan name e.g. Gold Plan 2026"),
        ("plan_type",        "ENUM, DEFAULT silver",       "bronze / silver / gold"),
        ("chit_type",        "VARCHAR(20), DEFAULT lucky_draw", "lucky_draw / auction"),
        ("monthly_amount",   "FLOAT, NOT NULL",            "Monthly installment amount (Rs.)"),
        ("total_members",    "INT, NOT NULL",              "Maximum members allowed"),
        ("current_members",  "INT, DEFAULT 0",             "Currently enrolled members"),
        ("duration_months",  "INT, NOT NULL",              "Total scheme duration in months"),
        ("total_value",      "FLOAT, NOT NULL",            "Total pool = monthly x members"),
        ("commission_rate",  "FLOAT, DEFAULT 5.0",         "Foreman commission percentage"),
        ("net_prize_amount", "FLOAT, NOT NULL",            "Prize after commission deduction"),
        ("start_date",       "DATE, NOT NULL",             "Scheme start date"),
        ("end_date",         "DATE, NOT NULL",             "Auto-calculated scheme end date"),
        ("status",           "ENUM, DEFAULT draft",        "draft / active / completed / cancelled"),
        ("created_by",       "INT, FK -> users.id",        "Admin who created the plan"),
        ("created_at",       "TIMESTAMP, Server Default",  "Plan creation timestamp"),
    ]),
    ("Table 4: plan_members", [
        ("id",             "INT, PK",                    "Unique enrollment ID"),
        ("user_id",        "INT, FK -> users.id",        "Enrolled member"),
        ("plan_id",        "INT, FK -> chit_plans.id",   "Plan joined"),
        ("ticket_number",  "VARCHAR(50), NULL",           "Member ticket e.g. TCK-1024-1"),
        ("joined_date",    "DATE, NOT NULL",              "Enrollment date"),
        ("has_won",        "BOOLEAN, DEFAULT FALSE",      "Whether member has won a draw"),
        ("is_active",      "BOOLEAN, DEFAULT TRUE",       "Enrollment active status"),
        ("reason",         "VARCHAR(50), DEFAULT personal", "Purpose of joining"),
        ("priority_score", "INT, DEFAULT 5",              "AI Emergency Priority (1=High, 5=Low)"),
        ("created_at",     "TIMESTAMP, Server Default",   "Enrollment timestamp"),
    ]),
    ("Table 5: monthly_draws", [
        ("id",                  "INT, PK",                       "Unique draw ID"),
        ("plan_id",             "INT, FK -> chit_plans.id",      "Chit plan this draw belongs to"),
        ("month_number",        "INT, NOT NULL",                  "Draw month number (1, 2 ... N)"),
        ("draw_date",           "DATE, NOT NULL",                 "Date of auction draw"),
        ("chit_value",          "FLOAT, NOT NULL",                "Total pool for this month"),
        ("winning_bid_amount",  "FLOAT, NULL",                    "Winning bid amount (Rs.)"),
        ("discount",            "FLOAT, DEFAULT 0.0",             "Pool minus winning bid"),
        ("company_commission",  "FLOAT, DEFAULT 0.0",             "5% of discount"),
        ("net_discount",        "FLOAT, DEFAULT 0.0",             "Discount after commission"),
        ("dividend_per_member", "FLOAT, DEFAULT 0.0",             "Net discount / total members"),
        ("net_payable",         "FLOAT, DEFAULT 0.0",             "Monthly due minus dividend"),
        ("prize_amount",        "FLOAT, NOT NULL",                "Amount winner actually receives"),
        ("winning_ticket_id",   "INT, FK -> plan_members.id",    "Winning member ticket"),
        ("winner_id",           "INT, FK -> users.id",           "Winning member"),
        ("status",              "ENUM, DEFAULT scheduled",        "scheduled / completed / cancelled"),
        ("conducted_by",        "INT, FK -> users.id",           "Admin who conducted the draw"),
        ("created_at",          "TIMESTAMP, Server Default",      "Draw record timestamp"),
    ]),
    ("Table 6: installments", [
        ("id",               "INT, PK",                      "Unique installment ID"),
        ("plan_member_id",   "INT, FK -> plan_members.id",   "Member enrollment"),
        ("plan_id",          "INT, FK -> chit_plans.id",     "Associated chit plan"),
        ("month_number",     "INT, NOT NULL",                 "Month serial number"),
        ("due_date",         "DATE, NOT NULL",                "Payment due date"),
        ("base_amount",      "FLOAT, NOT NULL",               "Base installment amount"),
        ("penalty_amount",   "FLOAT, DEFAULT 0.0",            "Late payment penalty"),
        ("total_amount",     "FLOAT, NOT NULL",               "Base + Penalty"),
        ("status",           "ENUM, DEFAULT pending",         "pending / paid / verified / overdue"),
        ("paid_date",        "DATE, NULL",                    "Actual payment date"),
        ("payment_method",   "VARCHAR(100), NULL",            "cash / bank_transfer / UPI"),
        ("transaction_id",   "VARCHAR(100), NULL",            "UPI / UTR reference number"),
        ("receipt_number",   "VARCHAR(50), UNIQUE, NULL",     "Auto-generated REC-2026-XXXX"),
        ("receipt_path",     "VARCHAR(255), NULL",            "PDF receipt file path"),
        ("verified_by",      "INT, FK -> users.id",          "Admin who verified payment"),
        ("verified_at",      "TIMESTAMP, NULL",               "Verification timestamp"),
        ("created_at",       "TIMESTAMP, Server Default",     "Record creation timestamp"),
    ]),
    ("Table 7: payments", [
        ("id",              "INT, PK",                       "Unique payment ID"),
        ("installment_id",  "INT, FK -> installments.id",    "Installment being paid"),
        ("user_id",         "INT, FK -> users.id",           "Member making the payment"),
        ("amount",          "FLOAT, NOT NULL",               "Payment amount (Rs.)"),
        ("payment_method",  "VARCHAR(50), NOT NULL",         "UPI / Cash / Bank Transfer"),
        ("transaction_id",  "VARCHAR(100), NULL",            "12-digit UTR reference"),
        ("status",          "ENUM, DEFAULT pending",         "pending / success / failed"),
        ("notes",           "TEXT, NULL",                    "Payment remarks"),
        ("created_at",      "TIMESTAMP, Server Default",     "Payment timestamp"),
    ]),
    ("Table 8: promissory_notes", [
        ("id",                "INT, PK",                      "Unique pronote record ID"),
        ("pronote_id",        "VARCHAR(50), UNIQUE, NOT NULL","Auto-generated PRN-2026-XXXX"),
        ("member_id",         "INT, FK -> users.id",          "Winner member"),
        ("chit_id",           "INT, FK -> chit_plans.id",     "Associated chit plan"),
        ("loan_amount",       "FLOAT, NOT NULL",               "Chit prize amount received"),
        ("issue_date",        "DATE, NOT NULL",                "Pronote issue date"),
        ("due_date",          "DATE, NOT NULL",                "Repayment due date"),
        ("interest_rate",     "FLOAT, DEFAULT 0.0",            "Applicable interest rate (%)"),
        ("repayment_amount",  "FLOAT, NOT NULL",               "Total repayment obligation"),
        ("purpose",           "TEXT, NULL",                    "Purpose stated by member"),
        ("status",            "ENUM, NOT NULL",                "pending_signature / active / paid / overdue"),
        ("signature",         "TEXT, NULL",                    "Base64 encoded Canvas digital signature"),
        ("created_at",        "TIMESTAMP, Server Default",     "Pronote generation timestamp"),
    ]),
    ("Table 9: guarantors", [
        ("id",                       "INT, PK",                  "Unique guarantor ID"),
        ("member_id",                "INT, FK -> users.id",      "Member this guarantor backs"),
        ("name",                     "VARCHAR(100), NOT NULL",   "Guarantor full name"),
        ("phone",                    "VARCHAR(15), NOT NULL",    "Guarantor phone number"),
        ("address",                  "TEXT, NOT NULL",           "Guarantor address"),
        ("relationship_to_member",   "VARCHAR(50), NOT NULL",    "Father / Spouse / Friend etc."),
        ("aadhar_number",            "VARCHAR(20), NOT NULL",    "Guarantor Aadhaar number"),
        ("document_path",            "VARCHAR(255), NULL",       "Uploaded guarantor document path"),
        ("created_at",               "TIMESTAMP, Server Default","Record creation timestamp"),
    ]),
    ("Table 10: notifications", [
        ("id",                 "INT, PK",                   "Unique notification ID"),
        ("user_id",            "INT, FK -> users.id",       "Recipient user"),
        ("title",              "VARCHAR(200), NOT NULL",    "Notification heading"),
        ("message",            "TEXT, NOT NULL",            "Full notification body"),
        ("notification_type",  "ENUM, DEFAULT info",        "success / warning / alert / info"),
        ("is_read",            "BOOLEAN, DEFAULT FALSE",    "Read / Unread status"),
        ("created_at",         "TIMESTAMP, Server Default", "Dispatch timestamp"),
    ]),
    ("Table 11: certificates", [
        ("id",                  "INT, PK",                    "Unique certificate ID"),
        ("user_id",             "INT, FK -> users.id",        "Certificate holder"),
        ("plan_id",             "INT, FK -> chit_plans.id",   "Completed scheme"),
        ("certificate_number",  "VARCHAR(50), UNIQUE",        "Certificate serial number"),
        ("file_path",           "VARCHAR(255), NULL",         "PDF certificate file path"),
        ("issued_by",           "INT, FK -> users.id",        "Admin who issued the certificate"),
        ("issued_at",           "TIMESTAMP, Server Default",  "Issue timestamp"),
    ]),
    ("Table 12: reports", [
        ("id",             "INT, PK",                  "Unique report ID"),
        ("report_type",    "ENUM, NOT NULL",            "members / kyc_status / installments / overdue / auctions / revenue / audit_log"),
        ("report_format",  "ENUM, NOT NULL",            "pdf / csv / excel"),
        ("file_path",      "VARCHAR(255), NULL",        "Generated report file path"),
        ("generated_by",   "INT, FK -> users.id",      "Admin who generated the report"),
        ("generated_at",   "TIMESTAMP, Server Default", "Report generation timestamp"),
    ]),
    ("Table 13: audit_logs", [
        ("id",           "INT, PK",                  "Unique log entry ID"),
        ("user_id",      "INT, FK -> users.id",      "User who triggered the action"),
        ("action",       "VARCHAR(100), NOT NULL",   "Action type e.g. LOGIN / PAYMENT_VERIFIED"),
        ("entity_type",  "VARCHAR(50), NULL",        "Affected entity e.g. installment"),
        ("entity_id",    "INT, NULL",                "ID of the affected record"),
        ("details",      "TEXT, NULL",               "Action description and metadata"),
        ("ip_address",   "VARCHAR(45), NULL",        "User IP address"),
        ("created_at",   "TIMESTAMP, Server Default","Log entry timestamp"),
    ]),
    ("Table 14: system_settings", [
        ("id",          "INT, PK",                   "Setting record ID"),
        ("key",         "VARCHAR(100), UNIQUE, NOT NULL", "Setting key e.g. commission_rate"),
        ("value",       "TEXT, NOT NULL",             "Setting value e.g. 5.0"),
        ("description", "TEXT, NULL",                 "Human-readable description"),
        ("updated_at",  "TIMESTAMP, Auto-update",     "Last update timestamp"),
        ("created_at",  "TIMESTAMP, Server Default",  "Creation timestamp"),
    ]),
    ("Table 15: password_resets", [
        ("id",           "INT, PK",                   "Unique token record ID"),
        ("user_id",      "INT, FK -> users.id",       "Account requesting reset"),
        ("reset_token",  "VARCHAR(255), UNIQUE, NOT NULL", "Hashed one-time reset token"),
        ("expires_at",   "TIMESTAMP, NOT NULL",        "Token expiry time"),
        ("is_used",      "BOOLEAN, DEFAULT FALSE",     "Whether token has been consumed"),
        ("created_at",   "TIMESTAMP, Server Default",  "Token generation timestamp"),
    ]),
]

# ── Find DB Design section: locate heading and remove old tables/paras ────────
body = doc.element.body
all_elems = list(body)

# Find start: paragraph with "4.2 DATABASE DESIGN"
# Find end: paragraph with "4.3 MODULE DESCRIPTION"
start_idx = None
end_idx   = None

for i, elem in enumerate(all_elems):
    tag = elem.tag.split('}')[-1]
    if tag == 'p':
        text = ''.join(r.text or '' for r in elem.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t'))
        if '4.2 DATABASE DESIGN' in text and start_idx is None:
            start_idx = i
            print(f"Found DB Design start at elem {i}: {repr(text[:60])}")
        if '4.3 MODULE DESCRIPTION' in text and start_idx is not None and end_idx is None:
            end_idx = i
            print(f"Found Module Desc end at elem {i}: {repr(text[:60])}")
            break

if start_idx is None or end_idx is None:
    print(f"start={start_idx}, end={end_idx} — checking all text...")
    for i, elem in enumerate(all_elems):
        tag = elem.tag.split('}')[-1]
        if tag == 'p':
            text = ''.join(r.text or '' for r in elem.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t'))
            if text.strip():
                print(f"  {i}: {repr(text.strip()[:80])}")
    exit()

# Remove all elements between start (inclusive) and end (exclusive)
elems_to_remove = all_elems[start_idx:end_idx]
for elem in elems_to_remove:
    body.remove(elem)
print(f"Removed {len(elems_to_remove)} elements from DB Design section")

# Re-fetch end index after removal (end_idx shifted)
all_elems = list(body)
new_end_idx = None
for i, elem in enumerate(all_elems):
    tag = elem.tag.split('}')[-1]
    if tag == 'p':
        text = ''.join(r.text or '' for r in elem.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t'))
        if '4.3 MODULE DESCRIPTION' in text:
            new_end_idx = i
            print(f"New end index: {i}")
            break

anchor_elem = all_elems[new_end_idx]

# ── Insert new section heading ─────────────────────────────────────────────────
def insert_before(new_elem, ref_elem):
    ref_elem.getparent().insert(list(ref_elem.getparent()).index(ref_elem), new_elem)

def new_para(text, bold=False, size=13, italic=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    r.font.size = Pt(size)
    return p

# Blank line
bp = new_para("")
bp.runs[0].font.size = Pt(6)
insert_before(bp._element, anchor_elem)

# "4.2 DATABASE DESIGN" heading
dhp = new_para("4.2 DATABASE DESIGN", bold=True, size=13)
insert_before(dhp._element, anchor_elem)

bp2 = new_para("")
bp2.runs[0].font.size = Pt(6)
insert_before(bp2._element, anchor_elem)

# ── Insert all 15 tables ──────────────────────────────────────────────────────
for tbl_heading, rows_data in ALL_TABLES:
    # Table heading para
    hp = new_para(tbl_heading, bold=True, italic=True, size=11)
    insert_before(hp._element, anchor_elem)

    # Build table
    tbl = doc.add_table(rows=1 + len(rows_data), cols=3)
    tbl.style = 'Table Grid'
    widths = [Cm(3.5), Cm(4.8), Cm(6.0)]

    # Header
    hdr = tbl.rows[0].cells
    for ci, col_name in enumerate(["Field", "Type", "Description"]):
        hdr[ci].width = widths[ci]
        set_cell_bg(hdr[ci], 'BFBFBF')
        set_row_height(tbl.rows[0], 1.0)
        p = hdr[ci].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(col_name)
        r.bold = True
        r.font.size = Pt(10)

    # Data rows
    for ri, (field, ftype, desc) in enumerate(rows_data):
        row = tbl.rows[ri + 1]
        set_row_height(row, 0.85)
        for ci, val in enumerate([field, ftype, desc]):
            row.cells[ci].width = widths[ci]
            p = row.cells[ci].paragraphs[0]
            r = p.add_run(val)
            r.font.size = Pt(10)

    insert_before(tbl._element, anchor_elem)

    # Blank line after table
    sep = new_para("")
    sep.runs[0].font.size = Pt(6)
    insert_before(sep._element, anchor_elem)

    print(f"Inserted: {tbl_heading}")

doc.save(out)
print(f"\nSaved: {out}")
shutil.copy(out, r'C:\Users\nares\.gemini\antigravity\brain\f74afe5d-2f73-4013-bc1e-3562e05df0f5\SecureChit_Report_With_DFDs.docx')
print("Copied to artifact dir!")
