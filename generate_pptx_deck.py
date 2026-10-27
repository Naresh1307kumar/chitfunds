import sys
import os

try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
except ImportError:
    os.system("pip install python-pptx")
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6]

    # Colors
    NAVY = RGBColor(26, 43, 73)      # #1A2B49
    GOLD = RGBColor(217, 119, 6)     # Amber-600
    DARK = RGBColor(15, 23, 42)      # Slate-900
    WHITE = RGBColor(255, 255, 255)
    LIGHT_BG = RGBColor(248, 250, 252)

    def add_header(slide, title_text, category_text="SECURECHIT SYSTEM"):
        # Header box
        tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(1.0))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p0 = tf.paragraphs[0]
        p0.text = category_text.upper()
        p0.font.size = Pt(11)
        p0.font.bold = True
        p0.font.color.rgb = GOLD
        
        p1 = tf.add_paragraph()
        p1.text = title_text
        p1.font.size = Pt(24)
        p1.font.bold = True
        p1.font.color.rgb = NAVY

    # SLIDE 1: Title
    s1 = prs.slides.add_slide(blank_layout)
    tb1 = s1.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(11.3), Inches(3.5))
    tf1 = tb1.text_frame
    p = tf1.paragraphs[0]
    p.text = "SECURE CHIT FUND MANAGEMENT SYSTEM"
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = NAVY

    p_sub = tf1.add_paragraph()
    p_sub.text = "Comprehensive System Architecture, Database Schemas & Diagrammatic Specifications"
    p_sub.font.size = Pt(18)
    p_sub.font.color.rgb = GOLD

    p_tech = tf1.add_paragraph()
    p_tech.text = "\nTech Stack: FastAPI (Python) | React.js (Vite) | PostgreSQL | TailwindCSS | Mermaid"
    p_tech.font.size = Pt(14)
    p_tech.font.color.rgb = DARK

    # SLIDE 2: Database Schema
    s2 = prs.slides.add_slide(blank_layout)
    add_header(s2, "Relational Database Table Schemas", "DATA ARCHITECTURE")
    tb2 = s2.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.7), Inches(5.2))
    tf2 = tb2.text_frame
    tf2.word_wrap = True

    tables_text = [
        ("users", "User Account & Role-Based Auth (id, username, email, password_hash, role, status, kyc_status)"),
        ("chit_plans", "Chit Scheme Specifications (id, plan_name, monthly_amount, total_members, chit_value, commission_rate)"),
        ("plan_members", "Plan Memberships & Ticket Numbers (id, user_id, plan_id, ticket_number, has_won, priority_score)"),
        ("monthly_draws", "Monthly Auction & Dividend Calculations (id, plan_id, winning_bid_amount, discount, company_commission, dividend_per_member, net_payable)"),
        ("installments", "Monthly Member Dues (id, plan_member_id, month_number, base_amount, total_amount, status)"),
        ("payments", "Transaction Records (id, installment_id, amount, payment_method, transaction_id, status)")
    ]

    for tbl, desc in tables_text:
        p = tf2.add_paragraph()
        p.text = f"• {tbl.upper()}: {desc}"
        p.font.size = Pt(14)
        p.font.bold = False
        p.font.color.rgb = DARK

    # SLIDE 3: ER Diagram
    s3 = prs.slides.add_slide(blank_layout)
    add_header(s3, "Entity-Relationship (ER) Diagram", "SYSTEM DESIGN")
    tb3 = s3.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.7), Inches(5.2))
    tf3 = tb3.text_frame
    tf3.word_wrap = True

    er_text = [
        "Primary & Foreign Key Relationships & Data Cardinality:",
        "1. USERS (1) ───< (N) PLAN_MEMBERS [User enrolls in multiple Chit Schemes]",
        "2. CHIT_PLANS (1) ───< (N) PLAN_MEMBERS [Scheme contains multiple enrolled member tickets]",
        "3. CHIT_PLANS (1) ───< (N) MONTHLY_DRAWS [Scheme conducts 1 Auction Draw per month]",
        "4. PLAN_MEMBERS (1) ───< (N) INSTALLMENTS [Member pays monthly installment dues]",
        "5. INSTALLMENTS (1) ───< (N) PAYMENTS [Installment generates online UPI payment records]",
        "6. MONTHLY_DRAWS (1) ─── (1) USERS [Winner user is awarded auction payout]"
    ]
    for line in er_text:
        p = tf3.add_paragraph()
        p.text = line
        p.font.size = Pt(14)
        p.font.color.rgb = DARK

    # SLIDE 4: DFD Diagram
    s4 = prs.slides.add_slide(blank_layout)
    add_header(s4, "Data Flow Diagrams (DFD Level 0 & Level 1)", "PROCESS FLOW")
    tb4 = s4.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.7), Inches(5.2))
    tf4 = tb4.text_frame
    tf4.word_wrap = True

    dfd_text = [
        "DFD Level 0 (Context Level):",
        "  • Member Entity ──> Register, Submit KYC, Join Plan, Pay Net Monthly Dues ──> System",
        "  • System ──> Digital Receipts, Dividend Breakdown, Monthly Schedule ──> Member Entity",
        "  • Admin Entity ──> Create Plans, Approve KYC, Conduct Monthly Auction ──> System",
        "  • System ──> AI Recommended Winners, System Audit Logs, Financial Reports ──> Admin Entity",
        "",
        "DFD Level 1 (Subsystem Level):",
        "  • 1.0 Auth & KYC Subsystem ──> Stores Profiles & Verification Docs in D1 Database",
        "  • 2.0 Plan & Member Mgmt Subsystem ──> Manages Schemes & Ticket Assignments in D2/D3 Database",
        "  • 3.0 Payment Gateway Subsystem ──> Updates Monthly Dues Status in D5 Database",
        "  • 4.0 AI Allocation & Auction Engine ──> Computes Dividends & Stores Draws in D4 Database"
    ]
    for line in dfd_text:
        p = tf4.add_paragraph()
        p.text = line
        p.font.size = Pt(13)
        p.font.color.rgb = DARK

    # SLIDE 5: Use Case Diagram
    s5 = prs.slides.add_slide(blank_layout)
    add_header(s5, "Use Case Design & Actor Roles", "USER INTERACTION")
    tb5 = s5.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.7), Inches(5.2))
    tf5 = tb5.text_frame
    tf5.word_wrap = True

    uc_text = [
        "Member Actor Use Cases:",
        "  • UC-1: Register Account & Login with Role Security",
        "  • UC-2: Upload Digital KYC Documents & Guarantor Profile",
        "  • UC-3: Explore & Join Available Auction Savings Schemes",
        "  • UC-4: View Scheme Schedule (Unlocked Current Due & Locked Future Months)",
        "  • UC-5: Make Online Monthly Dues Payment via Instant UPI",
        "  • UC-6: View & Print Official Payment Receipts",
        "",
        "Admin Actor Use Cases:",
        "  • UC-7: Review & Approve Member KYC Verification Requests",
        "  • UC-8: Create, Edit & Launch Chit Fund Savings Schemes",
        "  • UC-9: Conduct Monthly Auction Draw with Live Calculation Preview",
        "  • UC-10: Utilize AI Emergency Ranking Engine (Hospital > Education > Marriage)",
        "  • UC-11: Generate & Download System Financial Reports"
    ]
    for line in uc_text:
        p = tf5.add_paragraph()
        p.text = line
        p.font.size = Pt(13)
        p.font.color.rgb = DARK

    output_path = r"c:\chit fund\SecureChit_Presentation.pptx"
    prs.save(output_path)
    print(f"[SUCCESS] PowerPoint presentation saved to {output_path}")

if __name__ == "__main__":
    create_presentation()
