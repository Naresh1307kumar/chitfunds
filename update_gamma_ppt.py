from pptx import Presentation
from pptx.util import Pt
import sys

def move_slide(prs, old_index, new_index):
    xml_slides = prs.slides._sldIdLst  
    slides = list(xml_slides)
    xml_slides.remove(slides[old_index])
    xml_slides.insert(new_index, slides[old_index])

def add_bullet_slide(prs, title_text, bullet_points):
    try:
        slide_layout = prs.slide_layouts[1] # standard title+content layout
    except:
        slide_layout = prs.slide_layouts[0] # fallback

    slide = prs.slides.add_slide(slide_layout)
    
    if slide.shapes.title:
        slide.shapes.title.text = title_text
    
    try:
        tf = slide.shapes.placeholders[1].text_frame
        for i, point in enumerate(bullet_points):
            if i == 0:
                tf.text = point
            else:
                p = tf.add_paragraph()
                p.text = point
                p.level = 0
                
        # Adjust font size
        for paragraph in tf.paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(20)
    except:
        pass
        
    return slide

def main():
    ppt_path = r"C:\chit fund\Secure-Chit -demo-ppt.pptx"
    out_path = r"C:\chit fund\Secure-Chit -demo-ppt_updated.pptx"
    
    try:
        prs = Presentation(ppt_path)
    except Exception as e:
        print(f"Error opening PPT: {e}")
        sys.exit(1)
        
    initial_slide_count = len(prs.slides)
    
    # Slide A: User Workflow
    user_bullets = [
        "1. Registration & KYC: User signs up and submits documents for admin approval.",
        "2. Plan Enrollment: Selects a chit plan and pays the first installment via Razorpay.",
        "3. Live Bidding: Participates in the monthly reverse auction or requests AI emergency allocation.",
        "4. Digital Signature: Winner electronically signs the auto-generated Promissory Note using HTML5 Canvas.",
        "5. Prize Disbursal: Admin verifies the signature and releases the prize amount.",
        "6. Dividend & Dues: User enjoys automated monthly dividends and pays ongoing installments online."
    ]
    add_bullet_slide(prs, "User (Member) Workflow", user_bullets)
    
    # Slide B: Admin Workflow
    admin_bullets = [
        "1. Member & KYC Verification: Reviews new user registrations and verifies Aadhaar/PAN.",
        "2. Chit Group Creation: Creates new chit plans, sets rules, and schedules auction dates.",
        "3. Auction Monitoring: Oversees the live reverse auction and declares the winner.",
        "4. Dividend Calculation: System handles the 5% commission, Admin ensures dividends are distributed.",
        "5. Signature Verification: Admin reviews and approves the winner's digital signature on the Promissory Note.",
        "6. Prize Disbursal: Releases the prize money to the verified winner.",
        "7. Ledger Audit: Monitors monthly UPI payments and maintains the database audit logs."
    ]
    add_bullet_slide(prs, "Admin (Foreman) Workflow", admin_bullets)
    
    # Slide C: Integrated Workflow
    integrated_bullets = [
        "Onboarding Phase: User submits KYC -> Admin approves and creates groups.",
        "Auction Phase: User bids in real-time -> Admin monitors and System declares winner.",
        "Documentation Phase: System generates Pronote -> User signs digitally -> Admin verifies signature.",
        "Settlement Phase: Admin disburses prize -> System distributes dividends -> User pays ongoing dues."
    ]
    add_bullet_slide(prs, "Integrated System Workflow", integrated_bullets)
    
    # Reorder slides
    # We want to place the 3 new slides right after slide 2 (index 1). So they will be at index 2, 3, 4.
    for i in range(3):
        old_idx = initial_slide_count + i
        new_idx = 2 + i
        move_slide(prs, old_idx, new_idx)
    
    prs.save(out_path)
    print(f"SUCCESS: {out_path}")

if __name__ == "__main__":
    main()
