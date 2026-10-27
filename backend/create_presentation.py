import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def build_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Colors
    c_navy = RGBColor(15, 76, 129)      # #0F4C81
    c_dark_navy = RGBColor(26, 43, 73)  # #1A2B49
    c_gold = RGBColor(217, 119, 6)     # Amber/Gold
    c_bg_slate = RGBColor(248, 250, 252) # #F8FAFC
    c_white = RGBColor(255, 255, 255)
    c_dark_text = RGBColor(30, 41, 59)
    c_gray_text = RGBColor(100, 116, 139)
    c_light_blue = RGBColor(239, 246, 255)

    blank_layout = prs.slide_layouts[6]
    art_dir = r"C:\Users\nares\.gemini\antigravity\brain\f3eb3d7a-3368-4f7d-ad6c-f8c0de697e83"

    def add_header(slide, title_text, category="SECURE CHIT FUND MANAGEMENT SYSTEM"):
        banner = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(1.1))
        banner.fill.solid()
        banner.fill.fore_color.rgb = c_dark_navy
        banner.line.fill.background()

        tx_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.12), Inches(11), Inches(0.3))
        tf = tx_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = category.upper()
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = c_gold

        tx_box2 = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11), Inches(0.6))
        tf2 = tx_box2.text_frame
        tf2.word_wrap = True
        p2 = tf2.paragraphs[0]
        p2.text = title_text
        p2.font.size = Pt(24)
        p2.font.bold = True
        p2.font.color.rgb = c_white

    # -------------------------------------------------------------
    # SLIDE 1: Title Slide
    # -------------------------------------------------------------
    s1 = prs.slides.add_slide(blank_layout)
    bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = c_dark_navy
    bg1.line.fill.background()

    tbox = s1.shapes.add_textbox(Inches(1.5), Inches(2.2), Inches(10.333), Inches(3.5))
    tf = tbox.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Secure Chit Fund Management System"
    p.font.size = Pt(40)
    p.font.bold = True
    p.font.color.rgb = c_gold
    p.alignment = PP_ALIGN.CENTER

    p2 = tf.add_paragraph()
    p2.text = "Project Implementation & User Interface Screen Flow Presentation"
    p2.font.size = Pt(20)
    p2.font.color.rgb = c_white
    p2.alignment = PP_ALIGN.CENTER
    p2.space_before = Pt(15)

    p3 = tf.add_paragraph()
    p3.text = "Tech Stack: FastAPI | React.js (Vite) | PostgreSQL | TailwindCSS"
    p3.font.size = Pt(14)
    p3.font.color.rgb = RGBColor(148, 163, 184)
    p3.alignment = PP_ALIGN.CENTER
    p3.space_before = Pt(30)

    # -------------------------------------------------------------
    # SLIDE 2: Problem Statement
    # -------------------------------------------------------------
    s2 = prs.slides.add_slide(blank_layout)
    add_header(s2, "Problem Statement")

    probs = [
        ("Manual Paper Register Risks", "Traditional chit funds rely on paper ledgers prone to loss, damage, and manual calculation errors."),
        ("Lack of Member Visibility", "Subscribers cannot view live bidding details, winner payouts, or dividend shares online."),
        ("Security & Trust Deficit", "Unverified subscriber enrollment and informal receipts cause disputes between members and organizers.")
    ]

    for i, (title, desc) in enumerate(probs):
        y_top = Inches(1.5 + i * 1.8)
        card = s2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), y_top, Inches(11.733), Inches(1.5))
        card.fill.solid()
        card.fill.fore_color.rgb = c_bg_slate
        card.line.color.rgb = RGBColor(226, 232, 240)

        tb = s2.shapes.add_textbox(Inches(1.1), y_top + Inches(0.2), Inches(11.1), Inches(1.1))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = f"⚠️ {title}"
        p.font.size = Pt(18)
        p.font.bold = True
        p.font.color.rgb = c_navy

        p_desc = tf.add_paragraph()
        p_desc.text = desc
        p_desc.font.size = Pt(14)
        p_desc.font.color.rgb = c_dark_text
        p_desc.space_before = Pt(8)

    # -------------------------------------------------------------
    # SLIDE 3: Existing System
    # -------------------------------------------------------------
    s3 = prs.slides.add_slide(blank_layout)
    add_header(s3, "Existing System Analysis")

    ex_points = [
        ("Paper-Based Ledger Maintenance", "All subscriber details, installment collections, and draw records are maintained manually in physical books."),
        ("Physical Auction Presence Needed", "Subscribers must travel to physical offices to participate in monthly auctions or draw announcements."),
        ("Manual Dividend & Interest Calculations", "Dividends and organizer commissions are calculated by hand, leading to accounting delays."),
        ("Informal Identity Verification", "Subscriber identity and guarantor verification are done verbally without digital document proof.")
    ]

    for i, (title, desc) in enumerate(ex_points):
        y_top = Inches(1.5 + i * 1.35)
        card = s3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), y_top, Inches(11.733), Inches(1.15))
        card.fill.solid()
        card.fill.fore_color.rgb = RGBColor(254, 242, 242)
        card.line.color.rgb = RGBColor(254, 202, 202)

        tb = s3.shapes.add_textbox(Inches(1.1), y_top + Inches(0.12), Inches(11.1), Inches(0.9))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = f"❌ {title}"
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = RGBColor(185, 28, 28)

        p_desc = tf.add_paragraph()
        p_desc.text = desc
        p_desc.font.size = Pt(13)
        p_desc.font.color.rgb = c_dark_text

    # -------------------------------------------------------------
    # SLIDE 4: Proposed System
    # -------------------------------------------------------------
    s4 = prs.slides.add_slide(blank_layout)
    add_header(s4, "Proposed System Architecture")

    prop_points = [
        ("Automated Relational Database (PostgreSQL)", "Stores all chit schemes, member tickets, and transaction records with audited soft-delete security."),
        ("Online Web Portals (Admin & Member)", "Enables subscribers and organizers to access schemes, submit applications, and track winners online."),
        ("Dual Scheme Support (Lucky Draw & Auction)", "Supports Lucky Draw schemes (fixed payout) and Auction schemes (bidding discount with dividend deduction)."),
        ("🤖 AI-Powered Smart Winner Allocation", "Automatically calculates emergency scores (Hospital=100pts, Education=80pts) & auto-selects the optimal auction winner.")
    ]

    for i, (title, desc) in enumerate(prop_points):
        y_top = Inches(1.5 + i * 1.35)
        card = s4.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), y_top, Inches(11.733), Inches(1.15))
        card.fill.solid()
        card.fill.fore_color.rgb = c_light_blue
        card.line.color.rgb = RGBColor(191, 219, 254)

        tb = s4.shapes.add_textbox(Inches(1.1), y_top + Inches(0.12), Inches(11.1), Inches(0.9))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = f"✨ {title}"
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = c_navy

        p_desc = tf.add_paragraph()
        p_desc.text = desc
        p_desc.font.size = Pt(13)
        p_desc.font.color.rgb = c_dark_text

    # -------------------------------------------------------------
    # SLIDE 5: Existing vs Proposed Comparison Table
    # -------------------------------------------------------------
    s5 = prs.slides.add_slide(blank_layout)
    add_header(s5, "System Comparison: Existing vs Proposed System")

    table_shape = s5.shapes.add_table(6, 3, Inches(0.8), Inches(1.5), Inches(11.733), Inches(5.2))
    table = table_shape.table
    table.columns[0].width = Inches(2.733)
    table.columns[1].width = Inches(4.5)
    table.columns[2].width = Inches(4.5)

    headers = ["Feature / Aspect", "Traditional Existing System", "Proposed Secure Chit Fund System"]
    for j, h in enumerate(headers):
        cell = table.cell(0, j)
        cell.fill.solid()
        cell.fill.fore_color.rgb = c_dark_navy
        p = cell.text_frame.paragraphs[0]
        p.text = h
        p.font.bold = True
        p.font.size = Pt(15)
        p.font.color.rgb = c_white

    rows_data = [
        ("Record Maintenance", "Paper ledgers & physical receipt books (prone to loss)", "Automated PostgreSQL Database with Soft-Delete Audit Trails"),
        ("Auction & Bidding", "Physical attendance required at organizer office", "Digital Bidding Portal + Interactive Roulette Spinner"),
        ("Emergency Priority", "Random drawing or informal organizer bias", "Ranked Priority Engine (Hospital 1st, Education 2nd, etc.)"),
        ("Dividend Allocation", "Manual calculation; delayed next installment reduction", "Automated real-time dividend calculation & auto-deduction"),
        ("Transparency & Access", "Zero member portal; rely on verbal updates", "Dedicated Member & Admin Dashboards with live winner tracking")
    ]

    for i, row in enumerate(rows_data):
        for j, val in enumerate(row):
            cell = table.cell(i+1, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = c_white if i % 2 == 0 else c_bg_slate
            p = cell.text_frame.paragraphs[0]
            p.text = val
            p.font.size = Pt(13)
            p.font.color.rgb = c_dark_text

    # -------------------------------------------------------------
    # SLIDE 6: Implemented Module List
    # -------------------------------------------------------------
    s6 = prs.slides.add_slide(blank_layout)
    add_header(s6, "Implemented Module List")

    modules = [
        ("1. Public Landing Page Module", "Presents active investment schemes, plan breakdown details, and scheme details modal."),
        ("2. Authentication & Login Module", "Role-based authentication (Admin & Member Login) with token-based password reset."),
        ("3. Admin Management Portal Module", "Comprehensive dashboard for scheme creation (Lucky Draw vs Auction), member approvals, and draw management."),
        ("4. Member KYC & Guarantor Module", "Upload and verification of Aadhar/PAN identity documents and guarantor contact details.")
    ]

    for i, (m_title, m_desc) in enumerate(modules):
        col = i % 2
        row = i // 2
        x = Inches(0.8 + col * 6.0)
        y = Inches(1.5 + row * 2.7)

        box = s6.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, Inches(5.7), Inches(2.4))
        box.fill.solid()
        box.fill.fore_color.rgb = c_light_blue
        box.line.color.rgb = RGBColor(191, 219, 254)

        tb = s6.shapes.add_textbox(x + Inches(0.2), y + Inches(0.2), Inches(5.3), Inches(2.0))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = m_title
        p.font.bold = True
        p.font.size = Pt(17)
        p.font.color.rgb = c_navy

        p2 = tf.add_paragraph()
        p2.text = m_desc
        p2.font.size = Pt(14)
        p2.font.color.rgb = c_dark_text
        p2.space_before = Pt(10)

    # -------------------------------------------------------------
    # SLIDE 7: Screen 1 - Landing Page (Image + Click Flow)
    # -------------------------------------------------------------
    s7 = prs.slides.add_slide(blank_layout)
    add_header(s7, "Screen 1: Public Landing Page & Click Flow")

    img_path1 = os.path.join(art_dir, "user_screen_landing.png")
    if os.path.exists(img_path1):
        s7.shapes.add_picture(img_path1, Inches(0.8), Inches(1.5), width=Inches(6.2))

    card7 = s7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.3), Inches(1.5), Inches(5.233), Inches(5.5))
    card7.fill.solid()
    card7.fill.fore_color.rgb = c_bg_slate
    card7.line.color.rgb = RGBColor(203, 213, 225)

    tb7 = s7.shapes.add_textbox(Inches(7.5), Inches(1.7), Inches(4.833), Inches(5.1))
    tf7 = tb7.text_frame
    tf7.word_wrap = True
    p = tf7.paragraphs[0]
    p.text = "🖱️ UI Element Click Flow Breakdown:"
    p.font.bold = True
    p.font.size = Pt(15)
    p.font.color.rgb = c_navy

    clicks1 = [
        ("Login Button (Top Right)", "Redirects user to /login page for Member and Admin login."),
        ("Plans Link", "Smooth scrolls down to active chit investment cards section."),
        ("Join Now -> Button", "Opens Join Plan modal or redirects user to available schemes."),
        ("View Plans Button", "Scrolls down directly to available chit plans grid.")
    ]
    for btn, act in clicks1:
        p_b = tf7.add_paragraph()
        p_b.text = f"• {btn}: {act}"
        p_b.font.size = Pt(12)
        p_b.font.color.rgb = c_dark_text
        p_b.space_before = Pt(10)

    # -------------------------------------------------------------
    # SLIDE 8: Screen 2 - Member Login (Image + Click Flow)
    # -------------------------------------------------------------
    s8 = prs.slides.add_slide(blank_layout)
    add_header(s8, "Screen 2: Member Login & Password Reset")

    img_path2 = os.path.join(art_dir, "user_screen_login.png")
    if os.path.exists(img_path2):
        s8.shapes.add_picture(img_path2, Inches(0.8), Inches(1.5), width=Inches(6.2))

    card8 = s8.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.3), Inches(1.5), Inches(5.233), Inches(5.5))
    card8.fill.solid()
    card8.fill.fore_color.rgb = c_light_blue
    card8.line.color.rgb = RGBColor(191, 219, 254)

    tb8 = s8.shapes.add_textbox(Inches(7.5), Inches(1.7), Inches(4.833), Inches(5.1))
    tf8 = tb8.text_frame
    tf8.word_wrap = True
    p = tf8.paragraphs[0]
    p.text = "🖱️ UI Element Click Flow Breakdown:"
    p.font.bold = True
    p.font.size = Pt(15)
    p.font.color.rgb = c_navy

    clicks2 = [
        ("Username / Email Input", "Accepts subscriber username or registered email address."),
        ("Eye Icon (👁️)", "Toggles password text visibility between hidden dots and clear text."),
        ("Sign In -> Button", "Authenticates user credentials & redirects to Member Dashboard."),
        ("Forgot Password? Link", "Launches password recovery workflow with automated token redirection.")
    ]
    for btn, act in clicks2:
        p_b = tf8.add_paragraph()
        p_b.text = f"• {btn}: {act}"
        p_b.font.size = Pt(12)
        p_b.font.color.rgb = c_dark_text
        p_b.space_before = Pt(10)

    # -------------------------------------------------------------
    # SLIDE 9: Screen 3 - Admin Dashboard (Image + Click Flow)
    # -------------------------------------------------------------
    s9 = prs.slides.add_slide(blank_layout)
    add_header(s9, "Screen 3: Admin Management Dashboard")

    img_path3 = os.path.join(art_dir, "user_screen_admin_dashboard.png")
    if os.path.exists(img_path3):
        s9.shapes.add_picture(img_path3, Inches(0.8), Inches(1.5), width=Inches(6.2))

    card9 = s9.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.3), Inches(1.5), Inches(5.233), Inches(5.5))
    card9.fill.solid()
    card9.fill.fore_color.rgb = c_bg_slate
    card9.line.color.rgb = RGBColor(203, 213, 225)

    tb9 = s9.shapes.add_textbox(Inches(7.5), Inches(1.7), Inches(4.833), Inches(5.1))
    tf9 = tb9.text_frame
    tf9.word_wrap = True
    p = tf9.paragraphs[0]
    p.text = "🖱️ UI Element Click Flow Breakdown:"
    p.font.bold = True
    p.font.size = Pt(15)
    p.font.color.rgb = c_navy

    clicks3 = [
        ("Stats Cards (Members/Plans/Draws)", "Displays live real-time system metrics and active operational counts."),
        ("Download Operations Reports", "Generates instant PDF/Excel exports for member status and collections history."),
        ("Pending Actions Prompt", "Clicking 'Approve 1 member' opens member approval workflow directly."),
        ("Logout Button", "Clears administrator JWT session and redirects to /login.")
    ]
    for btn, act in clicks3:
        p_b = tf9.add_paragraph()
        p_b.text = f"• {btn}: {act}"
        p_b.font.size = Pt(12)
        p_b.font.color.rgb = c_dark_text
        p_b.space_before = Pt(10)

    # -------------------------------------------------------------
    # SLIDE 10: Screen 4 - Member Dashboard (Image + Click Flow)
    # -------------------------------------------------------------
    s10 = prs.slides.add_slide(blank_layout)
    add_header(s10, "Screen 4: Member Dashboard & Winner Tracker")

    img_path4 = os.path.join(art_dir, "user_screen_member_dashboard.png")
    if os.path.exists(img_path4):
        s10.shapes.add_picture(img_path4, Inches(0.8), Inches(1.5), width=Inches(6.2))

    card10 = s10.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.3), Inches(1.5), Inches(5.233), Inches(5.5))
    card10.fill.solid()
    card10.fill.fore_color.rgb = c_light_blue
    card10.line.color.rgb = RGBColor(191, 219, 254)

    tb10 = s10.shapes.add_textbox(Inches(7.5), Inches(1.7), Inches(4.833), Inches(5.1))
    tf10 = tb10.text_frame
    tf10.word_wrap = True
    p = tf10.paragraphs[0]
    p.text = "🖱️ UI Element Click Flow Breakdown:"
    p.font.bold = True
    p.font.size = Pt(15)
    p.font.color.rgb = c_navy

    clicks4 = [
        ("My Schemes & Winner History", "Displays joined schemes and month-by-month draw winner results."),
        ("KYC Documents (Upload / View ->)", "Navigates directly to Member KYC Document Upload Page."),
        ("Guarantor Details (Add / View ->)", "Opens Guarantor contact registration form."),
        ("Member Status Badge", "Displays active profile identity and KYC Verified badge.")
    ]
    for btn, act in clicks4:
        p_b = tf10.add_paragraph()
        p_b.text = f"• {btn}: {act}"
        p_b.font.size = Pt(12)
        p_b.font.color.rgb = c_dark_text
        p_b.space_before = Pt(10)

    # -------------------------------------------------------------
    # SLIDE 11: Screen 5 - Member KYC Upload (Image + Click Flow)
    # -------------------------------------------------------------
    s11 = prs.slides.add_slide(blank_layout)
    add_header(s11, "Screen 5: Member KYC Document Upload Page")

    img_path5 = os.path.join(art_dir, "user_screen_kyc_verification.png")
    if os.path.exists(img_path5):
        s11.shapes.add_picture(img_path5, Inches(0.8), Inches(1.5), width=Inches(6.2))

    card11 = s11.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.3), Inches(1.5), Inches(5.233), Inches(5.5))
    card11.fill.solid()
    card11.fill.fore_color.rgb = c_bg_slate
    card11.line.color.rgb = RGBColor(203, 213, 225)

    tb11 = s11.shapes.add_textbox(Inches(7.5), Inches(1.7), Inches(4.833), Inches(5.1))
    tf11 = tb11.text_frame
    tf11.word_wrap = True
    p = tf11.paragraphs[0]
    p.text = "🖱️ UI Element Click Flow Breakdown:"
    p.font.bold = True
    p.font.size = Pt(15)
    p.font.color.rgb = c_navy

    clicks5 = [
        ("Aadhar Front/Back Dropzones", "Click to select and upload Aadhar Card image files."),
        ("PAN Card Dropzone", "Click to select and upload PAN Card document file."),
        ("Passport Photo Dropzone", "Click to select and upload passport-size photo."),
        ("Action Required Alert", "Displays mandatory document submission instructions.")
    ]
    for btn, act in clicks5:
        p_b = tf11.add_paragraph()
        p_b.text = f"• {btn}: {act}"
        p_b.font.size = Pt(12)
        p_b.font.color.rgb = c_dark_text
        p_b.space_before = Pt(10)

    # -------------------------------------------------------------
    # SLIDE 12: Summary of Deliverables
    # -------------------------------------------------------------
    s12 = prs.slides.add_slide(blank_layout)
    add_header(s12, "Summary of Implemented Deliverables")

    c_box1 = s12.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.5), Inches(11.733), Inches(5.2))
    c_box1.fill.solid()
    c_box1.fill.fore_color.rgb = c_dark_navy

    tb1 = s12.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(10.9), Inches(4.5))
    tf1 = tb1.text_frame
    tf1.word_wrap = True
    p = tf1.paragraphs[0]
    p.text = "🎉 Project Progress Summary"
    p.font.bold = True
    p.font.size = Pt(22)
    p.font.color.rgb = c_gold

    summary_points = [
        "1. Public Landing Page — Complete with investment scheme cards & plan details breakdown.",
        "2. Member & Admin Login — JWT security, password toggle & username/email password recovery.",
        "3. Admin Management Dashboard — Operations statistics, reports download & member approvals.",
        "4. Member Dashboard & Winner Tracker — My schemes tracking & monthly winner breakdown.",
        "5. KYC Verification Page — 4-grid dropzone upload system for identity document approval."
    ]
    for pt in summary_points:
        p_pt = tf1.add_paragraph()
        p_pt.text = pt
        p_pt.font.size = Pt(15)
        p_pt.font.color.rgb = c_white
        p_pt.space_before = Pt(16)

    # Save presentation
    path1 = r"C:\chit fund\SecureChit_Final_Presentation.pptx"
    prs.save(path1)
    print(f"Presentation saved successfully to: {path1}")

if __name__ == "__main__":
    build_presentation()
