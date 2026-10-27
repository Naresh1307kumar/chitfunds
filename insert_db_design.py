import shutil
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

src = r'C:\chit fund\SecureChit_Report_With_DFDs_v2.docx'
out = r'C:\chit fund\SecureChit_Report_With_DFDs_v3.docx'
shutil.copy(src, out)
doc = Document(out)

# Find last paragraph (Figure 4.4 caption) — insert after it
target_para = None
for i, para in enumerate(doc.paragraphs):
    if 'Figure 4.4' in para.text and 'Four-Tier' in para.text:
        target_para = para
        print(f"Found insertion point at para {i}: {repr(para.text[:80])}")
        break

if not target_para:
    print("Target not found!")
    exit()

parent = target_para._element.getparent()

# ─── Helper: add paragraph at position ────────────────────────────────────────
def add_para(text, bold=False, size=11, align=WD_ALIGN_PARAGRAPH.LEFT, insert_after_elem=None, idx_offset=0):
    p = doc.add_paragraph()
    p.alignment = align
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    # Insert at correct position
    _idx = list(parent).index(insert_after_elem) + 1 + idx_offset
    parent.insert(_idx, p._element)
    parent.remove(p._element)
    parent.insert(_idx, p._element)
    return p

def add_table_section(heading, columns, rows, insert_after_elem, idx_offset):
    """Add heading + table after given element at offset"""
    # Heading paragraph
    hp = doc.add_paragraph()
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = hp.add_run(heading)
    r.bold = True
    r.font.size = Pt(11)
    _idx = list(parent).index(insert_after_elem) + 1 + idx_offset
    parent.insert(_idx, hp._element)
    parent.remove(hp._element)
    parent.insert(_idx, hp._element)
    idx_offset += 1

    # Table
    tbl = doc.add_table(rows=1 + len(rows), cols=len(columns))
    tbl.style = 'Table Grid'
    # Header row
    hdr = tbl.rows[0].cells
    for ci, col in enumerate(columns):
        hdr[ci].text = col
        for run in hdr[ci].paragraphs[0].runs:
            run.bold = True
            run.font.size = Pt(9)
    # Data rows
    for ri, row_data in enumerate(rows):
        cells = tbl.rows[ri + 1].cells
        for ci, val in enumerate(row_data):
            cells[ci].text = val
            for run in cells[ci].paragraphs[0].runs:
                run.font.size = Pt(9)
    # Insert table
    _idx2 = list(parent).index(insert_after_elem) + 1 + idx_offset
    parent.insert(_idx2, tbl._element)
    parent.remove(tbl._element)
    parent.insert(_idx2, tbl._element)
    idx_offset += 1

    # Empty line after table
    ep = doc.add_paragraph()
    ep.add_run("").font.size = Pt(6)
    _idx3 = list(parent).index(insert_after_elem) + 1 + idx_offset
    parent.insert(_idx3, ep._element)
    parent.remove(ep._element)
    parent.insert(_idx3, ep._element)
    idx_offset += 1

    return idx_offset

# ─── All 15 tables ─────────────────────────────────────────────────────────────

COLS = ["Column", "Data Type", "Constraints", "Description"]

tables_data = [
    {
        "heading": "4.2 DATABASE DESIGN",
        "is_main_heading": True,
    },
    {
        "heading": "TABLE 1: users",
        "rows": [
            ("id","Integer","PK, Auto-Increment","Unique user identifier"),
            ("username","VARCHAR(50)","UNIQUE, NOT NULL","Login username"),
            ("name","VARCHAR(100)","NOT NULL","Full name"),
            ("email","VARCHAR(150)","NOT NULL, Indexed","Email address"),
            ("phone","VARCHAR(15)","NOT NULL","Contact number"),
            ("password_hash","VARCHAR(255)","NOT NULL","BCrypt hashed password"),
            ("role","ENUM","NOT NULL","member / admin"),
            ("status","ENUM","NOT NULL","pending / approved / active / suspended / rejected"),
            ("aadhar_number","VARCHAR(20)","NULL","Aadhaar card number"),
            ("pan_number","VARCHAR(20)","NULL","PAN card number"),
            ("date_of_birth","DATE","NULL","Date of birth"),
            ("occupation","VARCHAR(100)","NULL","Member occupation"),
            ("kyc_status","ENUM","NOT NULL","not_submitted / pending / verified / rejected"),
            ("failed_login_attempts","Integer","DEFAULT 0","Brute-force login counter"),
            ("last_login","TIMESTAMP","NULL","Last successful login timestamp"),
            ("is_deleted","BOOLEAN","DEFAULT FALSE","Soft delete flag"),
            ("created_at","TIMESTAMP","Server Default","Registration timestamp"),
            ("updated_at","TIMESTAMP","Auto-update","Last modification timestamp"),
        ],
    },
    {
        "heading": "TABLE 2: kyc_documents",
        "rows": [
            ("id","Integer","PK","Unique KYC record ID"),
            ("user_id","Integer","FK → users.id","Document owner"),
            ("document_type","ENUM","NOT NULL","aadhar_front / aadhar_back / pan_card / address_proof / passport_photo"),
            ("file_path","VARCHAR(255)","NOT NULL","Uploaded file storage path"),
            ("status","ENUM","DEFAULT pending","pending / verified / rejected"),
            ("rejection_reason","TEXT","NULL","Reason if rejected by admin"),
            ("verified_by","Integer","FK → users.id","Admin who verified the document"),
            ("verified_at","TIMESTAMP","NULL","Verification timestamp"),
            ("created_at","TIMESTAMP","Server Default","Upload timestamp"),
        ],
    },
    {
        "heading": "TABLE 3: chit_plans",
        "rows": [
            ("id","Integer","PK","Unique plan ID"),
            ("plan_name","VARCHAR(100)","NOT NULL","Plan name (e.g., Gold Plan 2026)"),
            ("plan_type","ENUM","DEFAULT silver","bronze / silver / gold"),
            ("chit_type","VARCHAR(20)","DEFAULT lucky_draw","lucky_draw / auction"),
            ("monthly_amount","FLOAT","NOT NULL","Monthly installment amount (Rs.)"),
            ("total_members","Integer","NOT NULL","Maximum members allowed"),
            ("current_members","Integer","DEFAULT 0","Currently enrolled members"),
            ("duration_months","Integer","NOT NULL","Total scheme duration in months"),
            ("total_value","FLOAT","NOT NULL","Total pool = monthly x members"),
            ("commission_rate","FLOAT","DEFAULT 5.0","Foreman commission percentage"),
            ("net_prize_amount","FLOAT","NOT NULL","Prize after commission deduction"),
            ("start_date","DATE","NOT NULL","Scheme start date"),
            ("end_date","DATE","NOT NULL","Auto-calculated scheme end date"),
            ("status","ENUM","DEFAULT draft","draft / active / completed / cancelled"),
            ("created_by","Integer","FK → users.id","Admin who created the plan"),
            ("created_at","TIMESTAMP","Server Default","Plan creation timestamp"),
        ],
    },
    {
        "heading": "TABLE 4: plan_members",
        "rows": [
            ("id","Integer","PK","Unique enrollment ID"),
            ("user_id","Integer","FK → users.id","Enrolled member"),
            ("plan_id","Integer","FK → chit_plans.id","Plan joined"),
            ("ticket_number","VARCHAR(50)","NULL","Member ticket e.g. TCK-1024-1"),
            ("joined_date","DATE","NOT NULL","Enrollment date"),
            ("has_won","BOOLEAN","DEFAULT FALSE","Whether member has won a draw"),
            ("is_active","BOOLEAN","DEFAULT TRUE","Enrollment active status"),
            ("reason","VARCHAR(50)","DEFAULT personal","Purpose of joining"),
            ("priority_score","Integer","DEFAULT 5","AI Emergency Priority (1=High, 5=Low)"),
            ("created_at","TIMESTAMP","Server Default","Enrollment timestamp"),
        ],
    },
    {
        "heading": "TABLE 5: monthly_draws",
        "rows": [
            ("id","Integer","PK","Unique draw ID"),
            ("plan_id","Integer","FK → chit_plans.id","Chit plan this draw belongs to"),
            ("month_number","Integer","NOT NULL","Draw month number (1, 2 ... N)"),
            ("draw_date","DATE","NOT NULL","Date of auction draw"),
            ("chit_value","FLOAT","NOT NULL","Total pool for this month"),
            ("winning_bid_amount","FLOAT","NULL","Winning bid amount (Rs.)"),
            ("discount","FLOAT","DEFAULT 0.0","Pool minus winning bid"),
            ("company_commission","FLOAT","DEFAULT 0.0","5% of discount"),
            ("net_discount","FLOAT","DEFAULT 0.0","Discount after commission"),
            ("dividend_per_member","FLOAT","DEFAULT 0.0","Net discount / total members"),
            ("net_payable","FLOAT","DEFAULT 0.0","Monthly due minus dividend"),
            ("prize_amount","FLOAT","NOT NULL","Amount winner actually receives"),
            ("winning_ticket_id","Integer","FK → plan_members.id","Winning member ticket"),
            ("winner_id","Integer","FK → users.id","Winning member"),
            ("status","ENUM","DEFAULT scheduled","scheduled / completed / cancelled"),
            ("conducted_by","Integer","FK → users.id","Admin who conducted the draw"),
            ("created_at","TIMESTAMP","Server Default","Draw record timestamp"),
        ],
    },
    {
        "heading": "TABLE 6: installments",
        "rows": [
            ("id","Integer","PK","Unique installment ID"),
            ("plan_member_id","Integer","FK → plan_members.id","Member enrollment"),
            ("plan_id","Integer","FK → chit_plans.id","Associated chit plan"),
            ("month_number","Integer","NOT NULL","Month serial number"),
            ("due_date","DATE","NOT NULL","Payment due date"),
            ("base_amount","FLOAT","NOT NULL","Base installment amount"),
            ("penalty_amount","FLOAT","DEFAULT 0.0","Late payment penalty"),
            ("total_amount","FLOAT","NOT NULL","Base + Penalty"),
            ("status","ENUM","DEFAULT pending","pending / paid / verified / overdue"),
            ("paid_date","DATE","NULL","Actual payment date"),
            ("payment_method","VARCHAR(100)","NULL","cash / bank_transfer / UPI"),
            ("transaction_id","VARCHAR(100)","NULL","UPI / UTR reference number"),
            ("receipt_number","VARCHAR(50)","UNIQUE, NULL","Auto-generated REC-2026-XXXX"),
            ("receipt_path","VARCHAR(255)","NULL","PDF receipt file path"),
            ("verified_by","Integer","FK → users.id","Admin who verified payment"),
            ("verified_at","TIMESTAMP","NULL","Verification timestamp"),
            ("created_at","TIMESTAMP","Server Default","Record creation timestamp"),
        ],
    },
    {
        "heading": "TABLE 7: payments",
        "rows": [
            ("id","Integer","PK","Unique payment ID"),
            ("installment_id","Integer","FK → installments.id","Installment being paid"),
            ("user_id","Integer","FK → users.id","Member making the payment"),
            ("amount","FLOAT","NOT NULL","Payment amount (Rs.)"),
            ("payment_method","VARCHAR(50)","NOT NULL","UPI / Cash / Bank Transfer"),
            ("transaction_id","VARCHAR(100)","NULL","12-digit UTR reference"),
            ("status","ENUM","DEFAULT pending","pending / success / failed"),
            ("notes","TEXT","NULL","Payment remarks"),
            ("created_at","TIMESTAMP","Server Default","Payment timestamp"),
        ],
    },
    {
        "heading": "TABLE 8: promissory_notes",
        "rows": [
            ("id","Integer","PK","Unique pronote record ID"),
            ("pronote_id","VARCHAR(50)","UNIQUE, NOT NULL","Auto-generated PRN-2026-XXXX"),
            ("member_id","Integer","FK → users.id","Winner member"),
            ("chit_id","Integer","FK → chit_plans.id","Associated chit plan"),
            ("loan_amount","FLOAT","NOT NULL","Chit prize amount received"),
            ("issue_date","DATE","NOT NULL","Pronote issue date"),
            ("due_date","DATE","NOT NULL","Repayment due date"),
            ("interest_rate","FLOAT","DEFAULT 0.0","Applicable interest rate (%)"),
            ("repayment_amount","FLOAT","NOT NULL","Total repayment obligation"),
            ("purpose","TEXT","NULL","Purpose stated by member"),
            ("status","ENUM","NOT NULL","pending_signature / active / paid / overdue"),
            ("signature","TEXT","NULL","Base64 encoded Canvas digital signature PNG"),
            ("created_at","TIMESTAMP","Server Default","Pronote generation timestamp"),
        ],
    },
    {
        "heading": "TABLE 9: guarantors",
        "rows": [
            ("id","Integer","PK","Unique guarantor ID"),
            ("member_id","Integer","FK → users.id","Member this guarantor backs"),
            ("name","VARCHAR(100)","NOT NULL","Guarantor full name"),
            ("phone","VARCHAR(15)","NOT NULL","Guarantor phone number"),
            ("address","TEXT","NOT NULL","Guarantor address"),
            ("relationship_to_member","VARCHAR(50)","NOT NULL","Father / Spouse / Friend etc."),
            ("aadhar_number","VARCHAR(20)","NOT NULL","Guarantor Aadhaar number"),
            ("document_path","VARCHAR(255)","NULL","Uploaded guarantor document path"),
            ("created_at","TIMESTAMP","Server Default","Record creation timestamp"),
        ],
    },
    {
        "heading": "TABLE 10: notifications",
        "rows": [
            ("id","Integer","PK","Unique notification ID"),
            ("user_id","Integer","FK → users.id","Recipient user"),
            ("title","VARCHAR(200)","NOT NULL","Notification heading"),
            ("message","TEXT","NOT NULL","Full notification body"),
            ("notification_type","ENUM","DEFAULT info","success / warning / alert / info"),
            ("is_read","BOOLEAN","DEFAULT FALSE","Read / Unread status"),
            ("created_at","TIMESTAMP","Server Default","Dispatch timestamp"),
        ],
    },
    {
        "heading": "TABLE 11: certificates",
        "rows": [
            ("id","Integer","PK","Unique certificate ID"),
            ("user_id","Integer","FK → users.id","Certificate holder"),
            ("plan_id","Integer","FK → chit_plans.id","Completed scheme"),
            ("certificate_number","VARCHAR(50)","UNIQUE","Certificate serial number"),
            ("file_path","VARCHAR(255)","NULL","PDF certificate file path"),
            ("issued_by","Integer","FK → users.id","Admin who issued the certificate"),
            ("issued_at","TIMESTAMP","Server Default","Issue timestamp"),
        ],
    },
    {
        "heading": "TABLE 12: reports",
        "rows": [
            ("id","Integer","PK","Unique report ID"),
            ("report_type","ENUM","NOT NULL","members / kyc_status / installments / overdue / auctions / revenue / audit_log"),
            ("report_format","ENUM","NOT NULL","pdf / csv / excel"),
            ("file_path","VARCHAR(255)","NULL","Generated report file path"),
            ("generated_by","Integer","FK → users.id","Admin who generated the report"),
            ("generated_at","TIMESTAMP","Server Default","Report generation timestamp"),
        ],
    },
    {
        "heading": "TABLE 13: audit_logs",
        "rows": [
            ("id","Integer","PK","Unique log entry ID"),
            ("user_id","Integer","FK → users.id","User who triggered the action"),
            ("action","VARCHAR(100)","NOT NULL","Action type e.g. LOGIN / PAYMENT_VERIFIED"),
            ("entity_type","VARCHAR(50)","NULL","Affected entity e.g. installment"),
            ("entity_id","Integer","NULL","ID of the affected record"),
            ("details","TEXT","NULL","Action description and metadata"),
            ("ip_address","VARCHAR(45)","NULL","User IP address"),
            ("created_at","TIMESTAMP","Server Default","Log entry timestamp"),
        ],
    },
    {
        "heading": "TABLE 14: system_settings",
        "rows": [
            ("id","Integer","PK","Setting record ID"),
            ("key","VARCHAR(100)","UNIQUE, NOT NULL","Setting key e.g. commission_rate"),
            ("value","TEXT","NOT NULL","Setting value e.g. 5.0"),
            ("description","TEXT","NULL","Human-readable description"),
            ("updated_at","TIMESTAMP","Auto-update","Last update timestamp"),
            ("created_at","TIMESTAMP","Server Default","Creation timestamp"),
        ],
    },
    {
        "heading": "TABLE 15: password_resets",
        "rows": [
            ("id","Integer","PK","Unique token record ID"),
            ("user_id","Integer","FK → users.id","Account requesting reset"),
            ("reset_token","VARCHAR(255)","UNIQUE, NOT NULL","Hashed one-time reset token"),
            ("expires_at","TIMESTAMP","NOT NULL","Token expiry time"),
            ("is_used","BOOLEAN","DEFAULT FALSE","Whether token has been consumed"),
            ("created_at","TIMESTAMP","Server Default","Token generation timestamp"),
        ],
    },
]

# ─── Insert all content ────────────────────────────────────────────────────────
anchor = target_para._element
offset = 0

for item in tables_data:
    if item.get("is_main_heading"):
        # Main section heading
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r2 = p.add_run(item["heading"])
        r2.bold = True
        r2.font.size = Pt(13)
        _idx = list(parent).index(anchor) + 1 + offset
        parent.insert(_idx, p._element)
        parent.remove(p._element)
        parent.insert(_idx, p._element)
        offset += 1

        # Blank line
        ep = doc.add_paragraph()
        ep.add_run("").font.size = Pt(6)
        _idx2 = list(parent).index(anchor) + 1 + offset
        parent.insert(_idx2, ep._element)
        parent.remove(ep._element)
        parent.insert(_idx2, ep._element)
        offset += 1
    else:
        # Sub-heading
        hp = doc.add_paragraph()
        hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        rh = hp.add_run(item["heading"])
        rh.bold = True
        rh.font.size = Pt(11)
        _idx = list(parent).index(anchor) + 1 + offset
        parent.insert(_idx, hp._element)
        parent.remove(hp._element)
        parent.insert(_idx, hp._element)
        offset += 1

        # Table
        tbl = doc.add_table(rows=1 + len(item["rows"]), cols=4)
        tbl.style = 'Table Grid'
        hdr_cells = tbl.rows[0].cells
        for ci, col in enumerate(COLS):
            hdr_cells[ci].text = col
            for run in hdr_cells[ci].paragraphs[0].runs:
                run.bold = True
                run.font.size = Pt(9)
        for ri, row_data in enumerate(item["rows"]):
            cells = tbl.rows[ri + 1].cells
            for ci, val in enumerate(row_data):
                cells[ci].text = val
                for run in cells[ci].paragraphs[0].runs:
                    run.font.size = Pt(9)

        _idx2 = list(parent).index(anchor) + 1 + offset
        parent.insert(_idx2, tbl._element)
        parent.remove(tbl._element)
        parent.insert(_idx2, tbl._element)
        offset += 1

        # Blank line after table
        ep = doc.add_paragraph()
        ep.add_run("").font.size = Pt(6)
        _idx3 = list(parent).index(anchor) + 1 + offset
        parent.insert(_idx3, ep._element)
        parent.remove(ep._element)
        parent.insert(_idx3, ep._element)
        offset += 1

doc.save(out)
print(f"Saved: {out}")

shutil.copy(out, r'C:\Users\nares\.gemini\antigravity\brain\f74afe5d-2f73-4013-bc1e-3562e05df0f5\SecureChit_Report_With_DFDs.docx')
print("Copied to artifact dir!")
