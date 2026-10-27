import shutil
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

# Load the already-updated document (the one with Level 0 and Level 1 already inserted)
src = r'C:\chit fund\SecureChit_Report_With_DFDs.docx'
doc = Document(src)

img_level2 = r'C:\chit fund\DFD_Level_2_Auction_Pronote.png'

# Find "Auction Bidding & Pronote Workflow (Level 2 DFD)" paragraph
target_para = None
for i, para in enumerate(doc.paragraphs):
    if 'Auction Bidding' in para.text and 'Level 2' in para.text:
        target_para = para
        print(f"Found target at para {i}: {repr(para.text[:80])}")
        break

if target_para:
    parent = target_para._element.getparent()
    idx = list(parent).index(target_para._element)

    # Image paragraph
    img_p = doc.add_paragraph()
    img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    img_p.add_run().add_picture(img_level2, width=Inches(4.5))
    parent.insert(idx + 1, img_p._element)
    parent.remove(img_p._element)
    parent.insert(idx + 1, img_p._element)

    # Caption paragraph
    cap_p = doc.add_paragraph()
    cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cap_p.add_run("Figure 4.5: Level 2 DFD – Auction Bidding & Pronote Workflow")
    r.bold = True
    r.font.size = Pt(11)
    parent.insert(idx + 2, cap_p._element)
    parent.remove(cap_p._element)
    parent.insert(idx + 2, cap_p._element)

    print("Inserted Level 2 DFD image successfully!")
else:
    print("Target paragraph not found!")

doc.save(src)
print(f"Saved: {src}")

# Copy to artifact dir too
shutil.copy(src, r'C:\Users\nares\.gemini\antigravity\brain\f74afe5d-2f73-4013-bc1e-3562e05df0f5\SecureChit_Report_With_DFDs.docx')
print("Copied to artifact dir!")
