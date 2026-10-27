from pptx import Presentation
from pptx.util import Pt
import os

def add_bullet_slide(prs, title_text, bullet_points):
    slide_layout = prs.slide_layouts[1] # Title and Content
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    title.text = title_text
    
    tf = slide.shapes.placeholders[1].text_frame
    for i, point in enumerate(bullet_points):
        if i == 0:
            tf.text = point
        else:
            p = tf.add_paragraph()
            p.text = point
            p.level = 0
            
    # Adjust font size for better fit
    for paragraph in tf.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(20) # slightly smaller to accommodate more text
            
    return slide

def main():
    prs = Presentation()

    # Slide 1: Title
    slide_layout = prs.slide_layouts[0] 
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = "Secure Chit Fund Management System"
    subtitle.text = "External Viva Examination\n\nTeam Members:\nR.NARESHKUMAR (24UCS08)\nAPR.JEGATHEESH VARSHAN (24UCS06)\n\nProject Guide:\nMrs. SM. VALLI B.E., M.C.A., M.Phil"

    title.text_frame.paragraphs[0].runs[0].font.size = Pt(44)
    for p in subtitle.text_frame.paragraphs:
        for run in p.runs:
            run.font.size = Pt(24)

    # Slide 2: Introduction Part 1
    intro_bullets_1 = [
        "Core Objective: Digitizes traditional paper-based chit funds to ensure complete transparency, ledger auditability, and legal compliance.",
        "End-to-End Automation: Seamlessly handles member onboarding, KYC validation, live reverse auction bidding, and automated dividend distribution.",
        "Secure Online Payments: Integrated Razorpay gateway enables members to pay monthly installments online and download instant Digital Receipts."
    ]
    add_bullet_slide(prs, "Brief Introduction", intro_bullets_1)

    # Slide 3: Introduction Part 2
    intro_bullets_2 = [
        "AI-Powered Smart Allocation: Evaluates emergency priorities to select auction winners fairly without manual bias.",
        "Digital Promissory Note Module: Auto-generates an official legal Pronote (PRNYYYY-XXXX) for the winner.",
        "HTML5 Canvas Signature: The Pronote is activated only after the winner signs interactively via a touch/mouse pad.",
        "Modern Tech Stack: Robust full-stack architecture using React 18, Tailwind CSS, Python FastAPI REST services, and PostgreSQL.",
        "Security & Access: Features Role-Based Access Control (RBAC) portals for secure Member and Admin workflows."
    ]
    add_bullet_slide(prs, "Key System Highlights", intro_bullets_2)

    # Slide 4: Workflow
    workflow_bullets = [
        "1. Registration & KYC: User signs up and admin approves.",
        "2. Plan Enrollment: User selects a plan and pays initial installment via Razorpay.",
        "3. Live Bidding: Participates in monthly auction (AI emergency priority applied).",
        "4. Digital Signature: Winner electronically signs the Promissory Note.",
        "5. Prize Disbursal: Admin verifies signature and releases funds.",
        "6. Ongoing Payments: Subsequent dues calculated with automated monthly dividends."
    ]
    add_bullet_slide(prs, "System Workflow", workflow_bullets)

    # Slide 5: Input Screens
    add_bullet_slide(prs, "Input Screens", [
        "[ Please insert User Registration / Member Dashboard screenshot here ]",
        "[ Please insert Live Auction / Bidding screenshot here ]"
    ])

    # Slide 6: Output Screens
    add_bullet_slide(prs, "Output Screens", [
        "[ Please insert generated Digital Promissory Note screenshot here ]",
        "[ Please insert Payment Receipt / Admin Audit screenshot here ]"
    ])

    # Slide 7: Conclusion
    conclusion_bullets = [
        "Eliminates inefficiencies and opacity of conventional paper-based chit funds.",
        "Ensures fair prize allocation via reverse-auction and AI Emergency Priority.",
        "Enforces the statutory 5% commission cap through automated dividend distribution.",
        "Passed all 25 system test cases confirming functional and security requirements.",
        "Successfully digitizes India's community savings instrument while maintaining trust."
    ]
    add_bullet_slide(prs, "Conclusion", conclusion_bullets)

    # Slide 8: Future Enhancements
    future_bullets = [
        "React Native Mobile App: Cross-platform access for payments and bidding.",
        "NACH/e-NACH Auto-Debit: Automated collection to reduce default risks.",
        "Blockchain Audit Trail: Immutable transaction records using Hyperledger Fabric.",
        "AI Member Credit Scoring: Risk-weighted plan eligibility based on payment history."
    ]
    add_bullet_slide(prs, "Future Enhancements", future_bullets)

    ppt_path = r"c:\chit fund\SecureChit_Viva_Presentation_v2.pptx"
    prs.save(ppt_path)

if __name__ == "__main__":
    main()
